"""Moteurs locaux. Les sorties sont publiées atomiquement sans écrasement."""
from __future__ import annotations
import csv
import gzip
import html
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import threading
import time
import uuid
import zipfile
from pathlib import Path
from typing import Callable
from .catalog import family, IMAGES, AUDIO

ROOT = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))
APP_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else ROOT
CREATE_NO_WINDOW = 0x08000000 if os.name == "nt" else 0


class Cancelled(Exception): pass
class ConversionError(Exception): pass
class NoChange(Exception):
    """A valid operation whose result would not improve the original."""


def office_path():
    candidates = [os.environ.get("ZESTE_LIBREOFFICE", ""), shutil.which("soffice") or "", str(APP_DIR / "tools/LibreOffice/program/soffice.exe"), str(ROOT / "tools/LibreOffice/program/soffice.exe"), r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    return next((str(Path(p)) for p in candidates if p and Path(p).is_file()), None)


def ffmpeg_path():
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def unique_path(parent: Path, name: str):
    p = parent / name
    counter = 2
    while p.exists():
        p = parent / f"{Path(name).stem} ({counter}){Path(name).suffix}"
        counter += 1
    return p


class Engine:
    def __init__(self, cancel=None, progress=None, output_dir=None):
        self.cancel = cancel or threading.Event()
        self.progress = progress or (lambda value, text: None)
        self.output_dir = Path(output_dir) if output_dir else None
        self.warnings = []

    def check(self):
        if self.cancel.is_set(): raise Cancelled("Opération annulée. Les originaux sont conservés.")

    def run_process(self, argv, timeout=1800, *, env=None, check_exit=True):
        self.check()
        # No shell, and no unread pipes that can deadlock on verbose media output.
        with tempfile.TemporaryFile() as stream:
            proc = subprocess.Popen([str(x) for x in argv], stdin=subprocess.DEVNULL, stdout=stream, stderr=stream, creationflags=CREATE_NO_WINDOW, env=env)
            start = time.monotonic()
            try:
                while proc.poll() is None:
                    self.check()
                    if time.monotonic() - start > timeout: raise ConversionError("Le traitement dépasse le délai autorisé.")
                    time.sleep(.08)
                stream.seek(0)
                log = stream.read().decode("utf-8", errors="replace")
                if check_exit and proc.returncode: raise ConversionError(log[-2500:] or f"Le moteur a quitté avec le code {proc.returncode}.")
                return log
            finally:
                if proc.poll() is None:
                    # LibreOffice's console launcher owns a soffice.bin child.
                    # Stop only this conversion tree when cancelled/timed out.
                    if os.name == "nt":
                        try:
                            subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW, timeout=10)
                        except (OSError, subprocess.TimeoutExpired): pass
                    if proc.poll() is None: proc.kill()
                    proc.wait()

    def ff(self, args):
        # A local container must not fetch remote playlist/segment references.
        local_args = []
        for arg in args:
            if arg == "-i": local_args += ["-protocol_whitelist", "file,pipe"]
            local_args.append(arg)
        return self.run_process([ffmpeg_path(), "-hide_banner", "-loglevel", "error", "-nostdin", "-y", *local_args])

    def media_info(self, src):
        # The bundled wheel includes ffmpeg, not ffprobe.
        text = self.run_process([ffmpeg_path(), "-hide_banner", "-nostdin", "-protocol_whitelist", "file,pipe", "-i", str(src)], timeout=30, check_exit=False)
        m = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", text)
        return {"duration": int(m[1])*3600 + int(m[2])*60 + float(m[3]) if m else None,
                "audio_streams": len(re.findall(r"^\s*Stream #\d+:\d+[^\n]*?: Audio:", text, re.M)),
                "video_streams": len(re.findall(r"^\s*Stream #\d+:\d+[^\n]*?: Video:", text, re.M))}

    def duration(self, src):
        value = self.media_info(src)["duration"]
        if value is None: raise ConversionError("Durée du média non détectée.")
        return value

    def publish(self, staging: Path, parent: Path):
        """Every staging item is a complete result; os.rename on Windows cannot overwrite."""
        results = []
        self.check()
        for item in sorted(staging.iterdir()):
            self.check()
            if item.is_file() and not item.stat().st_size: raise ConversionError("Le moteur a produit un fichier vide.")
            target = unique_path(parent, item.name)
            while True:
                try:
                    os.rename(item, target)
                    break
                except FileExistsError:
                    target = unique_path(parent, item.name)
            results.append(str(target))
        if not results: raise ConversionError("Aucun résultat n’a été produit.")
        return results

    def execute(self, paths, operation, options=None):
        paths = [Path(p) for p in paths]
        if not paths: raise ConversionError("Ajoute au moins un fichier.")
        options = dict(options or {})
        self.check()
        grouped = operation in {"collage", "images_pdf", "merge_pdf", "join_media"} or operation in {"to:zip", "to:tar", "to:gz", "to:7z", "create_zip"}
        batches = [paths] if grouped else [[p] for p in paths]
        outputs = []
        unchanged = 0
        for index, batch in enumerate(batches):
            self.check()
            self.progress(int(index * 100 / len(batches)), batch[0].name)
            try:
                if operation == "duplicate" and any(p.is_symlink() or p.is_junction() for p in batch):
                    raise ConversionError("La duplication des liens et jonctions n’est pas autorisée.")
                batch = [p.resolve(strict=True) for p in batch]
                allow_directories = operation in {"to:zip", "to:tar", "to:gz", "to:7z", "create_zip", "duplicate", "compress"}
                if any(not p.is_file() and not (allow_directories and p.is_dir()) for p in batch):
                    raise ConversionError("Sélectionne un fichier local valide.")
                parent = (self.output_dir or batch[0].parent).resolve()
                parent.mkdir(parents=True, exist_ok=True)
                # Staging beside output ensures atomic renames even on another volume.
                with tempfile.TemporaryDirectory(prefix=".zeste-", dir=parent) as work:
                    result_dir = Path(work) / "result"
                    result_dir.mkdir()
                    self._dispatch(batch, operation, options, result_dir)
                    outputs.extend(self.publish(result_dir, parent))
            except NoChange as exc:
                unchanged += 1
                self.warnings.append(str(exc))
            except Cancelled:
                raise
            except Exception as exc:
                if len(batches) == 1: raise
                self.warnings.append(f"{batch[0].name} : échec — {exc}")
        if not outputs and unchanged != len(batches): raise ConversionError("Aucun fichier du lot n’a pu être traité.\n" + "\n".join(self.warnings))
        self.progress(100, "Terminé")
        return outputs

    def _dispatch(self, paths, operation, o, dest):
        from .extras import dispatch_extra
        if dispatch_extra(self, paths, operation, o, dest): return
        src = paths[0]
        kind = family(src)
        if operation == "create_zip": return self.archive(paths, "zip", dest, unpack_archives=False)
        if operation in {"to:zip", "to:tar", "to:gz", "to:7z"}:
            return self.archive(paths, operation[3:], dest)
        if operation == "extract": return self.extract(src, dest / (src.stem + " — extrait"))
        if operation == "metadata": return self.metadata(src, dest)
        if operation == "strip_metadata": return self.strip_metadata(src, dest)
        if operation == "qr": return self.qr(src, dest)
        if operation in {"images_pdf", "collage"}: return self.image_group(paths, operation, o, dest)
        if operation in {"merge_pdf", "split_pdf", "reorder_pdf", "rotate_pdf", "redact_pdf"}: return self.pdf_tool(paths, operation, o, dest)
        if operation == "join_media": return self.join_media(paths, dest)
        if operation.startswith("to:"):
            ext = operation[3:]
            if kind == "image": return self.image_convert(src, ext, o, dest)
            if kind in {"audio", "video"}: return self.media(src, ext, "convert", o, dest)
            if kind == "pdf": return self.pdf_convert(src, ext, o, dest)
            if kind == "subtitle": return self.subtitle(src, ext, dest)
            if kind == "sheet": return self.sheet(src, ext, dest)
            if kind in {"document", "slides"}: return self.document(src, ext, dest)
        if operation == "compress" and kind == "pdf": return self.pdf_tool(paths, operation, o, dest)
        if operation == "compress" and kind not in {"image", "audio", "video", "pdf"}: return self.archive(paths, "zip", dest)
        if kind == "image": return self.image_convert(src, src.suffix.lower().lstrip(".") if src.suffix.lower() not in {".svg", ".gif"} else "png", o, dest, operation)
        if kind in {"audio", "video"}: return self.media(src, src.suffix.lower().lstrip("."), operation, o, dest)
        raise ConversionError("Cette opération n’est pas compatible avec ce fichier.")

    def open_image(self, src):
        from PIL import Image, ImageOps
        import pillow_heif
        pillow_heif.register_heif_opener()
        if src.suffix.lower() == ".svg":
            from PySide6.QtSvg import QSvgRenderer
            from PySide6.QtGui import QImage, QPainter
            from PySide6.QtCore import QBuffer, QIODevice
            renderer = QSvgRenderer(str(src))
            if not renderer.isValid(): raise ConversionError("SVG invalide.")
            size = renderer.defaultSize()
            if size.width() * size.height() > 80_000_000: raise ConversionError("Image trop grande.")
            img = QImage(size, QImage.Format.Format_ARGB32)
            img.fill(0)
            painter = QPainter(img)
            renderer.render(painter)
            painter.end()
            buf = QBuffer()
            buf.open(QIODevice.OpenModeFlag.ReadWrite)
            img.save(buf, "PNG")
            return Image.open(io.BytesIO(bytes(buf.data()))).convert("RGBA")
        with Image.open(src) as im:
            im.seek(0)
            return ImageOps.exif_transpose(im).copy()

    def save_image(self, im, path, quality=88):
        from PIL import Image
        ext = path.suffix.lower()
        if ext in {".jpg", ".jpeg", ".pdf", ".bmp"}:
            rgba = im.convert("RGBA")
            white = Image.new("RGB", im.size, "white")
            white.paste(rgba, mask=rgba.getchannel("A"))
            im = white
        fmt = {".jpg": "JPEG", ".jpeg": "JPEG", ".tif": "TIFF", ".heic": "HEIF", ".heif": "HEIF"}.get(ext, ext[1:].upper())
        kwargs = {"quality": int(quality)} if fmt in {"JPEG", "WEBP", "AVIF", "HEIF"} else {}
        if fmt == "PNG": kwargs = {"optimize": True}
        im.save(path, format=fmt, **kwargs)

    def image_convert(self, src, ext, o, dest, operation="convert"):
        from PIL import Image, ImageEnhance, ImageDraw, ImageFont
        im = self.open_image(src)
        if operation == "resize":
            width, height = int(o.get("width", 1600)), int(o.get("height", 0))
            if width < 1: raise ConversionError("La largeur doit être positive.")
            height = height or max(1, round(im.height * width / im.width))
            if width * height > 80_000_000: raise ConversionError("Dimensions trop grandes.")
            im = im.resize((width, height), Image.Resampling.LANCZOS)
        if operation == "rotate": im = im.rotate(-int(o.get("angle", 90)), expand=True)
        if operation in {"edit_image", "annotate_image", "crop_image", "redact_image"}:
            from .photo import render_photo
            im = render_photo(im, o)
            # Editing always exports a lossless PNG, including redactions.
            ext = "png"
        out = dest / f"{src.stem} — {operation if operation != 'convert' else ext}.{ext}"
        self.save_image(im, out, o.get("quality", 82 if operation == "compress" else 90))
        target = float(o.get("target_mb", 0)) * 1024 ** 2
        percent = float(o.get("percent", 0))
        if percent: target = src.stat().st_size * percent / 100
        if operation == "compress" and target:
            for quality in [75, 60, 45, 30, 20]:
                self.check()
                if out.stat().st_size <= target: break
                self.save_image(im, out, quality)
            if out.stat().st_size > target:
                self.warnings.append(f"{src.name} : poids cible non atteint. Dimensions conservées.")
        if operation == "compress" and out.stat().st_size >= src.stat().st_size:
            self.warnings.append(f"{src.name} : la copie n’est pas plus légère que l’original.")

    def image_group(self, paths, operation, o, dest):
        from PIL import Image, ImageOps
        if operation == "images_pdf":
            import pymupdf as fitz
            with fitz.open() as doc:
                for src in paths:
                    self.check()
                    im = self.open_image(src)
                    stream = io.BytesIO()
                    im.save(stream, "PNG")
                    page = doc.new_page(width=im.width * .75, height=im.height * .75)
                    page.insert_image(page.rect, stream=stream.getvalue())
                doc.save(dest / "Images — assemblées.pdf", deflate=True)
            return
        columns = max(1, int(o.get("columns", 2)))
        cell, gap = int(o.get('cell',640)), int(o.get('gap',20))
        radius = int(o.get('radius',0))
        if not 32<=cell<=4000 or not 0<=gap<=500 or not 0<=radius<=cell//2: raise ConversionError('Dimensions de collage invalides.')
        rows = math.ceil(len(paths) / columns)
        if (columns*cell+(columns+1)*gap)*(rows*cell+(rows+1)*gap)>80_000_000: raise ConversionError("Trop d’images pour ce collage.")
        canvas = Image.new("RGB", (columns * cell + (columns + 1) * gap, rows * cell + (rows + 1) * gap), o.get("background", "#ffffff"))
        for index, src in enumerate(paths):
            self.check()
            im = ImageOps.contain(self.open_image(src).convert("RGB"), (cell, cell))
            x = gap + (index % columns) * (cell + gap) + (cell - im.width) // 2
            y = gap + (index // columns) * (cell + gap) + (cell - im.height) // 2
            if radius:
                from PIL import ImageDraw
                mask=Image.new('L',im.size,0); ImageDraw.Draw(mask).rounded_rectangle((0,0,im.width-1,im.height-1),radius=radius,fill=255)
                canvas.paste(im,(x,y),mask)
            else: canvas.paste(im, (x, y))
        canvas.save(dest / "Collage.jpg", quality=92)

    @staticmethod
    def audio_codec(ext):
        return {"mp3": ["-c:a", "libmp3lame", "-b:a", "192k"], "m4a": ["-c:a", "aac", "-b:a", "192k"], "aac": ["-c:a", "aac"], "wav": ["-c:a", "pcm_s16le"], "flac": ["-c:a", "flac"], "ogg": ["-c:a", "libvorbis", "-q:a", "5"], "opus": ["-c:a", "libopus", "-b:a", "128k"], "aiff": ["-c:a", "pcm_s16be"], "aif": ["-c:a", "pcm_s16be"], "wma": ["-c:a", "wmav2", "-b:a", "192k"]}.get(ext, ["-c:a", "aac"])

    def media(self, src, ext, operation, o, dest):
        if operation == "split_media":
            duration = self.duration(src)
            if o.get('points'):
                try: points=[float(value.strip()) for value in str(o['points']).split(',')]
                except ValueError: raise ConversionError('Points de coupe : utilise des secondes séparées par des virgules.')
                if not points or len(points)>1000 or any(not math.isfinite(p) or not 0<p<duration for p in points): raise ConversionError('Points de coupe hors du média.')
                points=[0,*sorted(set(points)),duration]
                for i,(start,end) in enumerate(zip(points,points[1:])):
                    self.check(); self.media(src,ext,'trim',{'start':start,'duration':end-start,'suffix':f'partie {i+1:03}'},dest)
                return
            seconds = float(o.get("segment", 30))
            if not math.isfinite(seconds) or seconds <= 0: raise ConversionError("La durée d’un segment doit être positive.")
            if duration / seconds > 10000: raise ConversionError("Ce découpage produirait trop de fichiers. Choisis une durée supérieure.")
            for i in range(math.ceil(duration / seconds)):
                self.check()
                start = i * seconds
                self.media(src, ext, "trim", {"start": start, "duration": min(seconds, duration-start), "suffix": f"partie {i+1:03}"}, dest)
            return
        if operation == "snapshot":
            if o.get('times'):
                try: times=[float(value.strip()) for value in str(o['times']).split(',')]
                except ValueError: raise ConversionError('Positions : utilise des secondes séparées par des virgules.')
                duration=self.duration(src)
                if len(times)>1000 or any(not math.isfinite(t) or not 0<=t<duration for t in times): raise ConversionError('Position hors de la vidéo.')
                for i,position in enumerate(times):
                    self.check(); self.ff(['-ss',str(position),'-i',src,'-frames:v','1',dest/f'{src.stem} — image {i+1:03}.png'])
                return
            if not math.isfinite(float(o.get('start',0))) or not 0<=float(o.get('start',0))<self.duration(src): raise ConversionError('Position hors de la vidéo.')
            self.ff(["-ss", str(o.get("start", 0)), "-i", src, "-frames:v", "1", dest / f"{src.stem} — image.png"])
            return
        suffix = str(o.get('suffix', ext if operation == 'convert' else operation))
        if re.search(r'[\\/:<>"|?*\x00-\x1f]', suffix) or not re.fullmatch(r"[a-z0-9]+", ext):
            raise ConversionError("Nom de sortie invalide.")
        out = dest / f"{src.stem} — {suffix}.{ext}"
        args = ["-i", src]
        if operation == "trim":
            if float(o.get("start", 0)) < 0 or float(o.get("duration", 10)) <= 0: raise ConversionError("Durées invalides.")
            args += ["-ss", str(o.get("start", 0)), "-t", str(o.get("duration", 10))]
        vf, af = [], []
        video = ext not in AUDIO and ext != "gif"
        if ext == "gif":
            args += ["-filter_complex", "fps=12,scale=640:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse", "-an", out]
            self.ff(args)
            return
        if operation == "speed":
            speed = float(o.get("speed", 1.5))
            if not .25 <= speed <= 4: raise ConversionError("Vitesse autorisée : 0,25 à 4.")
            if video: vf.append(f"setpts=PTS/{speed}")
            factors, remaining = [], speed
            while remaining > 2: factors.append(2); remaining /= 2
            while remaining < .5: factors.append(.5); remaining /= .5
            factors.append(remaining)
            af += [f"atempo={x}" for x in factors]
        if operation == "crop_video":
            width, height = int(o.get("width", 640)), int(o.get("height", 480))
            x, y = int(o.get("x", 0)), int(o.get("y", 0))
            if min(width, height) <= 0 or min(x, y) < 0: raise ConversionError("Zone de recadrage invalide.")
            vf.append(f"crop={width}:{height}:{x}:{y}")
        if operation == "normalize":
            loudness,peak,span=float(o.get('loudness',-16)),float(o.get('true_peak',-1.5)),float(o.get('loudness_range',11))
            if not (-50<=loudness<=-5 and -9<=peak<=0 and 1<=span<=50): raise ConversionError('Réglages de normalisation invalides.')
            af.append(f'loudnorm=I={loudness}:TP={peak}:LRA={span}')
        if operation == 'channels':
            left,right=float(o.get('left_gain',1)),float(o.get('right_gain',1))
            if not (0<=left<=3 and 0<=right<=3): raise ConversionError('Volumes de canaux invalides.')
            af += ['aformat=channel_layouts=stereo',f'pan=stereo|c0={left}*c0|c1={right}*c1']
        if operation == "bleep":
            start, end = float(o.get("start", 0)), float(o.get("end", 1))
            if start < 0 or end <= start: raise ConversionError("Intervalle invalide.")
            # Keep the original duration; mute only the selected span, overlay a sine.
            filt = f"[0:a]volume=enable='between(t,{start},{end})':volume=0[a];sine=frequency=1000:sample_rate=48000:duration={end-start},volume=0.12,adelay={round(start*1000)}:all=1[b];[a][b]amix=inputs=2:duration=first:normalize=0[out]"
            args += ["-filter_complex", filt, "-map", "[out]"]
        if video:
            if ext == "webm": codecs = ["-c:v", "libvpx-vp9", "-crf", "34", "-b:v", "0", "-c:a", "libopus"]
            elif ext == "wmv": codecs = ["-c:v", "wmv2", "-c:a", "wmav2"]
            elif ext == "avi": codecs = ["-c:v", "mpeg4", "-q:v", "4", "-c:a", "libmp3lame"]
            else: codecs = ["-c:v", "libx264", "-preset", "fast", "-crf", str(o.get("crf", 28 if operation == "compress" else 21)), "-c:a", "aac", "-b:a", "160k"]
            vf.append("scale=trunc(iw/2)*2:trunc(ih/2)*2")
            args += codecs
            if ext in {"mp4", "mov", "m4v"}: args += ["-movflags", "+faststart"]
        else: args += ["-vn", *self.audio_codec(ext)]
        target = float(o.get("target_mb", 0)) * 1024**2
        if o.get("percent"): target = src.stat().st_size * float(o["percent"]) / 100
        if not math.isfinite(target) or target < 0: raise ConversionError("Poids cible invalide.")
        if operation == "compress" and target:
            rate = int(target * 8 / self.duration(src) * .94)
            if video:
                if rate < 200_000: raise ConversionError("Poids cible trop petit pour cette durée. Choisis un poids supérieur.")
                # Bitrate mode must not retain CRF / quality mode from defaults.
                args = self.without_options(args, {"-crf", "-q:v", "-b:v", "-b:a"})
                args += ["-b:v", str(rate - 128_000), "-maxrate", str(rate - 128_000), "-bufsize", str(rate*2), "-b:a", "128k"]
            elif ext in {"mp3", "m4a", "aac", "ogg", "opus", "wma"}:
                args = self.without_options(args, {"-q:a", "-b:a"})
                # Vorbis rejects very high mono bitrates and low-rate source
                # combinations. Use a supported sample rate and conservative cap.
                if ext == "ogg": args += ["-ar", "48000"]
                args += ["-b:a", str(max(32000, min(rate, 160000 if ext == "ogg" else 320000)))]
        if vf: args += ["-vf", ",".join(vf)]
        if af: args += ["-af", ",".join(af)]
        if operation == "channels": args += ["-ac", str(int(o.get("channels", 1)))]
        if operation == "mute": args += ["-an"]
        args += ["-map_metadata", "-1", out]
        self.ff(args)
        if target and out.stat().st_size > target * 1.05: self.warnings.append(f"{src.name} : poids cible approximatif dépassé.")
        if operation == "compress" and out.stat().st_size >= src.stat().st_size: self.warnings.append(f"{src.name} : la copie n’est pas plus légère.")

    @staticmethod
    def without_options(args, names):
        result, skip = [], False
        for arg in args:
            if skip: skip = False
            elif arg in names: skip = True
            else: result.append(arg)
        return result

    def join_media(self, paths, dest):
        if len(paths) < 2: raise ConversionError("Sélectionne au moins deux médias.")
        kind = family(paths[0])
        if kind == 'video':
            from .extras import _video_size
            width,height = _video_size(self,paths[0]); width += width%2; height += height%2
        if any(family(p) != kind for p in paths): raise ConversionError("Assemble uniquement des vidéos ensemble ou des audios ensemble.")
        with tempfile.TemporaryDirectory(prefix="zeste-join-") as td:
            parts = []
            for i, path in enumerate(paths):
                part = Path(td) / f"part{i:04}.{'mp4' if kind == 'video' else 'wav'}"
                args = ["-i", path]
                if kind == "video":
                    # Uniform frame and audio characteristics, including silent sources.
                    has_audio = self.media_info(path)["audio_streams"] > 0
                    if not has_audio: args += ["-f", "lavfi", "-i", "anullsrc=r=48000:cl=stereo"]
                    args += ["-map", "0:v:0", "-map", "0:a:0" if has_audio else "1:a:0", "-vf", f"scale={width}:{height}:force_original_aspect_ratio=decrease,pad={width}:{height}:(ow-iw)/2:(oh-ih)/2,setsar=1,fps=30,setpts=PTS-STARTPTS", "-af", "aresample=48000,asetpts=PTS-STARTPTS,apad", "-c:v", "libx264", "-preset", "fast", "-pix_fmt", "yuv420p", "-c:a", "aac", "-ar", "48000", "-ac", "2", "-shortest"]
                else: args += ["-vn", "-ar", "48000", "-ac", "2", "-c:a", "pcm_s16le"]
                self.ff([*args, part])
                parts.append(part)
            listing = Path(td) / "concat.txt"
            listing.write_text("\n".join(f"file '{p.name}'" for p in parts), encoding="utf-8")
            self.ff(["-f", "concat", "-safe", "1", "-i", listing, "-map", "0:v:0?", "-map", "0:a:0", "-c", "copy", dest / f"Assemblage.{'mp4' if kind == 'video' else 'wav'}"])

    @staticmethod
    def pages(spec, count):
        result = []
        for bit in spec.split(","):
            if "-" in bit:
                a, b = map(int, bit.strip().split("-"))
                result.extend(range(a-1, b, 1) if b >= a else range(a-1, b-2, -1))
            else: result.append(int(bit.strip()) - 1)
        if not result or any(p < 0 or p >= count for p in result): raise ConversionError("Numéro de page hors du document.")
        return result

    def pdf_tool(self, paths, op, o, dest):
        import pymupdf as fitz
        src = paths[0]
        if op == "merge_pdf":
            with fitz.open() as result:
                for path in paths:
                    self.check()
                    with fitz.open(path) as part: result.insert_pdf(part)
                result.save(dest / "PDF — fusion.pdf", garbage=4, deflate=True)
            return
        with fitz.open(src) as doc:
            if doc.needs_pass: raise ConversionError("PDF protégé : déverrouille une copie avant de continuer.")
            if op == "split_pdf":
                for i in range(len(doc)):
                    self.check()
                    with fitz.open() as one:
                        one.insert_pdf(doc, from_page=i, to_page=i)
                        one.save(dest / f"{src.stem} — page {i+1:03}.pdf")
                return
            if op == "reorder_pdf" and "page_plan" in o:
                plan = o["page_plan"]
                if not plan or len(plan)>10000: raise ConversionError("Sélectionne au moins une page (10 000 maximum).")
                with fitz.open() as arranged:
                    for entry in plan:
                        index,rotation=int(entry['page']),int(entry.get('rotation',0))
                        if index<0 or index>=len(doc) or rotation%90: raise ConversionError("Page ou rotation invalide.")
                        self.check(); arranged.insert_pdf(doc, from_page=index,to_page=index)
                        arranged[-1].set_rotation((doc[index].rotation+rotation)%360)
                    arranged.set_metadata(doc.metadata)
                    arranged.save(dest/f"{src.stem} — pages organisées.pdf",garbage=4,deflate=True)
                return
            if op == "reorder_pdf": doc.select(self.pages(str(o.get("pages", "1")), len(doc)))
            if op == "rotate_pdf":
                for page in doc: page.set_rotation((page.rotation + int(o.get("angle", 90))) % 360)
            if op == "redact_pdf":
                text = str(o.get("text", "")).strip()
                if not text: raise ConversionError("Indique le texte à masquer.")
                found = 0
                for page in doc:
                    for rect in page.search_for(text):
                        page.add_redact_annot(rect, fill=(0, 0, 0))
                        found += 1
                    page.apply_redactions()
                if not found: raise ConversionError("Texte introuvable. Un PDF scanné nécessite un OCR préalable.")
            out = dest / f"{src.stem} — {op}.pdf"
            options = dict(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True)
            if op == "compress":
                # Preserve page content, forms, links and metadata. Packing PDF
                # objects avoids expanding files that already use object streams.
                options.update(use_objstms=1, compression_effort=100, no_new_id=True)
            self.check()
            doc.save(out, **options)
            if op == "compress" and out.stat().st_size >= src.stat().st_size:
                out.unlink()  # Only our private staging output, never the source.
                raise NoChange(f"{src.name} : déjà optimisé sans perte. Original conservé ; aucun nouveau fichier créé.")

    def pdf_convert(self, src, ext, o, dest):
        import pymupdf as fitz
        if ext == "docx":
            from pdf2docx import Converter
            converter = Converter(str(src))
            try: converter.convert(str(dest / (src.stem + ".docx")), multi_processing=False)
            finally: converter.close()
            return
        with fitz.open(src) as doc:
            if doc.needs_pass: raise ConversionError("PDF protégé.")
            if ext in {"png", "jpg"}:
                for i, page in enumerate(doc):
                    self.check()
                    page.get_pixmap(dpi=int(o.get("dpi", 200)), alpha=False).save(dest / f"{src.stem} — page {i+1:03}.{ext}")
            elif ext == "pptx":
                from pptx import Presentation
                from pptx.util import Pt
                prs = Presentation()
                prs.slide_width, prs.slide_height = Pt(doc[0].rect.width), Pt(doc[0].rect.height)
                for page in doc:
                    self.check()
                    stream = io.BytesIO(page.get_pixmap(dpi=150, alpha=False).tobytes("png"))
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    slide.shapes.add_picture(stream, 0, 0, width=prs.slide_width, height=prs.slide_height)
                prs.save(dest / (src.stem + ".pptx"))
                self.warnings.append("PowerPoint : chaque page est une image, les éléments ne sont pas éditables individuellement.")
            elif ext == "xlsx":
                from openpyxl import Workbook
                wb = Workbook()
                wb.remove(wb.active)
                tables = 0
                for index, page in enumerate(doc):
                    self.check()
                    for table in page.find_tables().tables:
                        tables += 1
                        ws = wb.create_sheet(f"Page {index+1} Tableau {tables}"[:31])
                        for row in table.extract(): ws.append([safe_cell(x) for x in row])
                if not tables: raise ConversionError("Aucun tableau détecté dans ce PDF. Le texte scanné nécessite un OCR.")
                wb.save(dest / (src.stem + ".xlsx"))
            else:
                text = "\n\n".join(page.get_text(sort=True) for page in doc)
                if not text.strip(): self.warnings.append("Aucun texte extractible : ce PDF semble scanné. Aucun OCR automatique n’a été appliqué.")
                if ext in {"txt", "md"}: (dest / f"{src.stem}.{ext}").write_text(text, encoding="utf-8")
                elif ext == "html": (dest / (src.stem + ".html")).write_text("<!doctype html><meta charset='utf-8'><pre>" + html.escape(text) + "</pre>", encoding="utf-8")
                elif ext == "epub": self.make_epub(text, src.stem, dest / (src.stem + ".epub"))
                else: raise ConversionError("Export PDF non pris en charge.")

    def make_epub(self, text, title, out):
        from ebooklib import epub
        book = epub.EpubBook()
        book.set_identifier(str(uuid.uuid4()))
        book.set_title(title)
        book.set_language("fr")
        chapter = epub.EpubHtml(title=title, file_name="content.xhtml", lang="fr")
        chapter.content = "<h1>" + html.escape(title) + "</h1>" + "".join("<p>"+html.escape(p)+"</p>" for p in text.splitlines())
        book.add_item(chapter)
        book.toc = (chapter,)
        book.spine = ["nav", chapter]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(str(out), book)

    def libreoffice(self, src, ext, dest):
        tool = office_path()
        if not tool: raise ConversionError("Cette conversion nécessite le moteur LibreOffice. Ouvre Réglages pour le configurer.")
        console_tool = Path(tool).with_suffix(".com")
        if os.name == "nt" and console_tool.is_file(): tool = str(console_tool)
        # Python selects a writable temporary root (and can fall back to cwd).
        # Keep this profile short: deeply nested destination paths break LO on
        # Windows. Its native scratch files must use this same writable root.
        with tempfile.TemporaryDirectory(prefix="zeste-office-") as td:
            profile = Path(td) / "profile"
            scratch, converted = Path(td) / "temp", Path(td) / "converted"
            scratch.mkdir()
            converted.mkdir()
            # Isolated profile: no first-run window, no recent documents or user macros.
            profile.mkdir()
            (profile / "user").mkdir()
            (profile / "user/registrymodifications.xcu").write_text('<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>', encoding="utf-8")
            env = {**os.environ, "TEMP": str(scratch), "TMP": str(scratch)}
            log = self.run_process([tool, "-env:UserInstallation="+profile.as_uri(), "--headless", "--nologo", "--nodefault", "--norestore", "--convert-to", ext, "--outdir", converted, src], 180, env=env)
            expected = converted / (src.stem + "." + ext.split(":")[0])
            if not expected.is_file() or not expected.stat().st_size:
                raise ConversionError("LibreOffice n’a pas pu convertir ce document ou ce format.\n" + log[-1500:])
            # Do not publish Office's temporary .~lock files as conversion results.
            shutil.copy2(expected, dest / expected.name)

    def document(self, src, ext, dest):
        if family(src) == "slides" and ext in {"png", "jpg"}:
            with tempfile.TemporaryDirectory() as td:
                self.libreoffice(src, "pdf", Path(td))
                return self.pdf_convert(Path(td) / (src.stem + ".pdf"), ext, {}, dest)
        if office_path() and (ext in {"pdf", "docx", "odt", "rtf", "pptx"} or src.suffix.lower() in {".doc", ".pages", ".ppt", ".pptx", ".odp"}):
            return self.libreoffice(src, ext, dest)
        import pypandoc
        out = dest / (src.stem + "." + ext)
        fmt = {"md": "markdown", "txt": "plain", "htm": "html"}.get(src.suffix.lower()[1:], src.suffix.lower()[1:])
        if fmt == "plain": fmt = "markdown"
        if ext == "pdf":
            import pymupdf as fitz
            markup = pypandoc.convert_file(str(src), "html", format=fmt, extra_args=["--standalone"])
            story = fitz.Story(html=markup)
            writer = fitz.DocumentWriter(str(out))
            page_rect = fitz.paper_rect("a4")
            more = True
            while more:
                self.check()
                dev = writer.begin_page(page_rect)
                more, _ = story.place(page_rect + (40, 40, -40, -40))
                story.draw(dev)
                writer.end_page()
            writer.close()
            self.warnings.append("PDF généré avec une mise en page simplifiée. LibreOffice conserve mieux la mise en page bureautique.")
        else:
            output_fmt = {"txt": "plain", "md": "markdown"}.get(ext, ext)
            pypandoc.convert_file(str(src), output_fmt, format=fmt, outputfile=str(out), extra_args=["--standalone"])

    def sheet(self, src, ext, dest):
        if ext in {"pdf", "ods"} or src.suffix.lower() in {".xls", ".ods"}: return self.libreoffice(src, ext, dest)
        from openpyxl import Workbook, load_workbook
        if src.suffix.lower() == ".xlsx":
            wb = load_workbook(src, data_only=True, read_only=True)
            sheets = [(ws.title, list(ws.values)) for ws in wb]
            wb.close()
        else:
            with src.open(encoding="utf-8-sig", newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try: dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error: dialect = csv.excel_tab if src.suffix.lower() == ".tsv" else csv.excel
                sheets = [(src.stem[:31], list(csv.reader(f, dialect)))]
        if ext == "xlsx":
            wb = Workbook()
            wb.remove(wb.active)
            for title, rows in sheets:
                ws = wb.create_sheet(title)
                for row in rows: ws.append([safe_cell(x) for x in row])
            wb.save(dest / (src.stem + ".xlsx"))
        else:
            for title, rows in sheets:
                safe_title = re.sub(r'[<>:"/\\|?*]', "_", title)
                name = src.stem + (" — " + safe_title if len(sheets) > 1 else "")
                with (dest / f"{name}.{ext}").open("w", encoding="utf-8-sig", newline="") as f:
                    csv.writer(f, delimiter="\t" if ext == "tsv" else ",").writerows([[safe_cell(x) for x in row] for row in rows])
            self.warnings.append("CSV/TSV conserve les valeurs, sans mise en forme ; une sortie par feuille.")

    def subtitle(self, src, ext, dest):
        text = src.read_text(encoding="utf-8-sig")
        if ext == "txt":
            text = re.sub(r"(?m)^WEBVTT.*$|^\d+\s*$|^.*-->.*$", "", text)
            text = re.sub(r"<[^>]+>", "", text)
            text = re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"
        elif ext == "vtt":
            text = "WEBVTT\n\n" + re.sub(r"(\d{2}:\d{2}:\d{2}),(\d{3})", r"\1.\2", text)
        elif ext == "srt":
            self.ff(["-i", src, dest / (src.stem + ".srt")])
            return
        (dest / f"{src.stem}.{ext}").write_text(text, encoding="utf-8")

    def archive(self, paths, ext, dest, unpack_archives=True):
        # Archive-to-archive means content conversion, not nesting the old archive.
        with tempfile.TemporaryDirectory(prefix="zeste-archive-", dir=dest.parent) as td:
            if unpack_archives and len(paths) == 1 and family(paths[0]) == "archive":
                unpack = Path(td) / "contents"
                self.extract(paths[0], unpack)
                inputs = list(unpack.iterdir())
            else: inputs = paths
            if not inputs: raise ConversionError("Archive vide.")
            names = [p.name.casefold() for p in inputs]
            if len(set(names)) != len(names): raise ConversionError("Plusieurs éléments portent le même nom. Renomme-les avant de les archiver.")
            name = paths[0].stem if len(paths) == 1 else "Fichiers"
            suffix = ".tar.gz" if ext == "gz" and (len(inputs) != 1 or inputs[0].is_dir()) else "."+ext
            out = dest / (name + suffix)
            if ext == "zip":
                with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=7) as z:
                    for path in inputs:
                        self.check()
                        if path.is_dir():
                            z.write(path, path.name + "/")
                            for item in path.rglob("*"):
                                self.check()
                                if item.is_symlink() or item.is_junction() or item.resolve().is_relative_to(dest.parent.resolve()): continue
                                z.write(item, Path(path.name) / item.relative_to(path))
                        else: z.write(path, path.name)
            elif ext in {"tar", "gz"}:
                if suffix == ".gz":
                    with inputs[0].open("rb") as f, gzip.open(out, "wb") as g: shutil.copyfileobj(f, g)
                else:
                    with tarfile.open(out, "w:gz" if ext == "gz" else "w") as t:
                        for path in inputs:
                            self.check()
                            t.add(path, arcname=path.name, filter=lambda info: None if info.issym() or info.islnk() else info)
            elif ext == "7z":
                import py7zr
                with py7zr.SevenZipFile(out, "w") as z:
                    for path in inputs:
                        self.check()
                        z.writeall(path, arcname=path.name)

    def extract(self, src, dest):
        dest.mkdir(parents=True, exist_ok=False)
        max_bytes, max_files = 8 * 1024**3, 30_000
        def validate(name):
            cleaned = name.replace("\\", "/")
            target = (dest / cleaned).resolve()
            if ":" in cleaned or not target.is_relative_to(dest.resolve()) or cleaned.startswith("/"):
                raise ConversionError("Archive refusée : chemin de sortie dangereux.")
            return target
        if zipfile.is_zipfile(src):
            with zipfile.ZipFile(src) as z:
                infos = z.infolist()
                if len(infos) > max_files or sum(i.file_size for i in infos) > max_bytes: raise ConversionError("Archive trop volumineuse pour l’extraction automatique.")
                for info in infos:
                    self.check()
                    target = validate(info.filename)
                    if (info.external_attr >> 16) & 0o170000 == 0o120000: raise ConversionError("Liens symboliques non autorisés dans l’archive.")
                    if info.is_dir(): target.mkdir(parents=True, exist_ok=True)
                    else:
                        target.parent.mkdir(parents=True, exist_ok=True)
                        with z.open(info) as f, target.open("xb") as out: shutil.copyfileobj(f, out)
        elif tarfile.is_tarfile(src):
            with tarfile.open(src) as t:
                infos = t.getmembers()
                if len(infos) > max_files or sum(i.size for i in infos) > max_bytes: raise ConversionError("Archive trop volumineuse.")
                for info in infos:
                    self.check()
                    validate(info.name)
                    if not (info.isfile() or info.isdir()): raise ConversionError("Archive contenant des liens ou fichiers spéciaux.")
                t.extractall(dest, filter="data")
        elif src.suffix.lower() == ".7z":
            import py7zr
            with py7zr.SevenZipFile(src) as z:
                infos = z.list()
                if len(infos) > max_files or sum(i.uncompressed or 0 for i in infos) > max_bytes: raise ConversionError("Archive trop volumineuse.")
                for i in infos:
                    validate(i.filename)
                    if getattr(i, "is_symlink", False): raise ConversionError("Liens non autorisés.")
                z.extractall(dest)
        elif src.suffix.lower() == ".rar":
            import rarfile
            seven = ROOT / "tools/7zip/7z.exe"
            if seven.exists(): rarfile.SEVENZIP_TOOL = str(seven)
            with rarfile.RarFile(src) as z:
                infos = z.infolist()
                if len(infos) > max_files or sum(i.file_size for i in infos) > max_bytes: raise ConversionError("Archive trop volumineuse.")
                for i in infos:
                    self.check()
                    target = validate(i.filename)
                    if i.is_symlink(): raise ConversionError("Liens non autorisés.")
                    if i.isdir(): target.mkdir(parents=True, exist_ok=True)
                    else:
                        target.parent.mkdir(parents=True, exist_ok=True)
                        with z.open(i) as f, target.open("xb") as out: shutil.copyfileobj(f, out)
        elif src.suffix.lower() == ".gz":
            with gzip.open(src, "rb") as f, (dest / src.stem).open("xb") as out:
                total = 0
                while chunk := f.read(1024 * 1024):
                    self.check()
                    total += len(chunk)
                    if total > max_bytes: raise ConversionError("Archive trop volumineuse.")
                    out.write(chunk)
        else: raise ConversionError("Format d’archive non pris en charge.")

    def metadata(self, src, dest):
        data = {"nom": src.name, "taille_octets": src.stat().st_size, "type": family(src)}
        kind = family(src)
        if kind == "image":
            from PIL import ExifTags
            im = self.open_image(src)
            data.update(dimensions=list(im.size), mode=im.mode, exif={ExifTags.TAGS.get(k, str(k)): str(v) for k, v in im.getexif().items()}, informations={k: str(v)[:1000] for k, v in im.info.items() if k not in {"exif", "icc_profile"}})
        elif kind == "pdf":
            import pymupdf as fitz
            with fitz.open(src) as doc: data.update(pages=len(doc), metadata=doc.metadata)
        elif kind in {"video", "audio"}:
            with tempfile.TemporaryDirectory() as td:
                path = Path(td) / "meta.txt"
                self.ff(["-i", src, "-f", "ffmetadata", path])
                data.update(duree_secondes=self.duration(src), metadata=path.read_text(encoding="utf-8"))
        elif zipfile.is_zipfile(src):
            with zipfile.ZipFile(src) as z: data["fichiers"] = [{"nom": i.filename, "taille": i.file_size} for i in z.infolist()]
        (dest / f"{src.stem} — informations.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    def strip_metadata(self, src, dest):
        kind = family(src)
        out = dest / f"{src.stem} — sans métadonnées{src.suffix}"
        if kind == "image":
            if src.suffix.lower() == ".svg": out = out.with_suffix(".png")
            im = self.open_image(src)
            im.info.clear()
            self.save_image(im, out, 95)
        elif kind == "pdf":
            import pymupdf as fitz
            with fitz.open(src) as doc:
                from .pdf_metadata import clear_pdf_metadata
                clear_pdf_metadata(doc)
                self.warnings.append(f"{src.name} : les mentions de logiciel et de droits sont conservées ; le contenu du PDF n’est pas anonymisé.")
                doc.save(out, garbage=4, deflate=True)
        elif kind in {"audio", "video"}: self.ff(["-i", src, "-map", "0", "-c", "copy", "-map_metadata", "-1", out])
        else: raise ConversionError("Nettoyage des métadonnées non pris en charge pour ce format.")

    def qr(self, src, dest):
        import cv2
        import numpy as np
        im = np.array(self.open_image(src).convert("RGB"))
        found, values, _, _ = cv2.QRCodeDetector().detectAndDecodeMulti(im)
        if not found or not any(values): raise ConversionError("Aucun QR code lisible détecté.")
        (dest / f"{src.stem} — QR.txt").write_text("\n".join(v for v in values if v), encoding="utf-8")


def safe_cell(value):
    # Exports from external text must not introduce spreadsheet formulas.
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
