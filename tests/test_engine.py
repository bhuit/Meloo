import hashlib
import json
import os
import threading
import zipfile
from pathlib import Path
import pytest
from PIL import Image, ImageDraw
import pymupdf as fitz
from zeste.engine import Engine, ConversionError, Cancelled, office_path


@pytest.fixture()
def image(tmp_path):
    path = tmp_path / "Photo été & soleil.png"
    im = Image.new("RGB", (320, 240), "#ff9d53")
    ImageDraw.Draw(im).rectangle((20, 20, 100, 100), fill="#537fbd")
    im.save(path)
    return path


@pytest.fixture()
def pdf(tmp_path):
    path = tmp_path / "rapport.pdf"
    with fitz.open() as doc:
        for i in range(3):
            page = doc.new_page()
            page.insert_text((50, 60), f"PAGE {i+1} / DOCUMENT CONFIDENTIEL / Zeste")
            page.insert_text((50, 100), "Un texte de test pour les conversions.")
        doc.set_metadata({"author": "Auteur test", "title": "Rapport test"})
        doc.save(path)
    return path


@pytest.fixture(scope="module")
def media(tmp_path_factory):
    root = tmp_path_factory.mktemp("media")
    video, audio = root / "video.mp4", root / "audio.wav"
    engine = Engine()
    engine.ff(["-f", "lavfi", "-i", "testsrc=size=320x240:rate=12:duration=2", "-f", "lavfi", "-i", "sine=frequency=440:duration=2", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", video])
    engine.ff(["-f", "lavfi", "-i", "sine=frequency=440:duration=2", audio])
    return video, audio


@pytest.mark.parametrize("ext", ["jpg", "webp", "avif", "heic", "tiff", "bmp", "gif", "pdf"])
def test_image_formats(image, ext, tmp_path):
    before = hashlib.sha256(image.read_bytes()).hexdigest()
    output = Path(Engine().execute([image], "to:" + ext)[0])
    assert output.exists() and output.stat().st_size > 0
    if ext == "pdf":
        with fitz.open(output) as doc: assert len(doc) == 1
    else:
        with Image.open(output) as im: assert im.size == (320, 240)
    assert hashlib.sha256(image.read_bytes()).hexdigest() == before


def test_no_overwrite_and_resize(image):
    first = Path(Engine().execute([image], "resize", {"width": 100})[0])
    data = first.read_bytes()
    second = Path(Engine().execute([image], "resize", {"width": 80})[0])
    assert first != second and first.read_bytes() == data
    assert Image.open(second).size == (80, 60)


def test_edit_redact_crop(image):
    out = Path(Engine().execute([image], "edit_image", {"crop": [0, 0, 200, 200], "annotations": [{"type": "redact", "box": [10, 10, 40, 40]}]})[0])
    with Image.open(out) as im:
        assert im.size == (200, 200)
        assert im.getpixel((20, 20))[:3] == (0, 0, 0)


@pytest.mark.parametrize("op", ["compress", "rotate", "metadata", "strip_metadata", "collage", "images_pdf"])
def test_image_tools(image, op):
    results = Engine().execute([image], op)
    assert all(Path(p).stat().st_size for p in results)


def test_qr(tmp_path):
    import qrcode
    source = tmp_path / "QR.png"
    qrcode.make("ZESTE_TEST_2026").save(source)
    result = Path(Engine().execute([source], "qr")[0])
    assert result.read_text() == "ZESTE_TEST_2026"


@pytest.mark.parametrize("ext", ["txt", "md", "png", "jpg", "pptx", "epub", "html", "docx"])
def test_pdf_conversion(pdf, ext):
    outputs = Engine().execute([pdf], "to:"+ext, {"dpi": 72})
    assert all(Path(p).stat().st_size for p in outputs)
    if ext in {"png", "jpg"}: assert len(outputs) == 3
    if ext in {"txt", "md"}: assert "PAGE 3" in Path(outputs[0]).read_text(encoding="utf-8")


def test_pdf_redaction_and_reorder(pdf):
    out = Path(Engine().execute([pdf], "redact_pdf", {"text": "CONFIDENTIEL"})[0])
    with fitz.open(out) as doc: assert all("CONFIDENTIEL" not in p.get_text() for p in doc)
    out = Path(Engine().execute([pdf], "reorder_pdf", {"pages": "3,1,2"})[0])
    with fitz.open(out) as doc: assert "PAGE 3" in doc[0].get_text()
    with fitz.open(pdf) as doc: assert "CONFIDENTIEL" in doc[0].get_text()


@pytest.mark.parametrize("op", ["compress", "split_pdf", "merge_pdf", "rotate_pdf", "metadata", "strip_metadata"])
def test_pdf_tools(pdf, op):
    results = Engine().execute([pdf, pdf] if op == "merge_pdf" else [pdf], op)
    assert all(Path(p).stat().st_size for p in results)


@pytest.mark.parametrize("ext", ["mp3", "m4a", "flac", "ogg", "opus", "aiff", "wma"])
def test_audio_formats(media, ext, tmp_path):
    out = Path(Engine(output_dir=tmp_path).execute([media[1]], "to:"+ext)[0])
    assert 1.8 <= Engine().duration(out) <= 2.2


@pytest.mark.parametrize("ext", ["mp4", "mov", "mkv", "webm", "avi", "wmv", "gif", "mp3"])
def test_video_formats(media, ext, tmp_path):
    out = Path(Engine(output_dir=tmp_path).execute([media[0]], "to:"+ext)[0])
    assert out.stat().st_size > 0
    if ext != "gif": assert Engine().duration(out) >= 1.8


@pytest.mark.parametrize("op,options", [("trim", {"start": .5, "duration": .7}), ("normalize", {}), ("bleep", {"start": .3, "end": .7}), ("channels", {"channels": 2}), ("speed", {"speed": 1.5}), ("split_media", {"segment": 1}), ("join_media", {}), ("compress", {"target_mb": .03}), ("metadata", {}), ("strip_metadata", {})])
def test_audio_tools(media, op, options, tmp_path):
    paths = [media[1], media[1]] if op == "join_media" else [media[1]]
    results = Engine(output_dir=tmp_path).execute(paths, op, options)
    assert all(Path(p).stat().st_size > 0 for p in results)


@pytest.mark.parametrize("op,options", [("snapshot", {}), ("crop_video", {"width": 160, "height": 120}), ("mute", {}), ("speed", {"speed": 2.5}), ("compress", {}), ("split_media", {"segment": 1}), ("join_media", {})])
def test_video_tools(media, op, options, tmp_path):
    paths = [media[0], media[0]] if op == "join_media" else [media[0]]
    results = Engine(output_dir=tmp_path).execute(paths, op, options)
    assert all(Path(p).stat().st_size > 0 for p in results)


@pytest.mark.parametrize("ext", ["zip", "tar", "gz", "7z"])
def test_archive_roundtrip(image, ext, tmp_path):
    source_hash = hashlib.sha256(image.read_bytes()).digest()
    archive = Engine().execute([image], "to:"+ext)[0]
    output = Path(Engine().execute([archive], "extract")[0])
    restored = next(p for p in output.rglob("*") if p.is_file())
    assert hashlib.sha256(restored.read_bytes()).digest() == source_hash


def test_archive_traversal_rejected(tmp_path):
    attack = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(attack, "w") as z: z.writestr("../escaped.txt", "unsafe")
    with pytest.raises(ConversionError): Engine().execute([attack], "extract")
    assert not (tmp_path / "escaped.txt").exists()
    assert list(tmp_path.iterdir()) == [attack]


def test_cancel_does_not_publish(image):
    cancellation = threading.Event()
    cancellation.set()
    with pytest.raises(Cancelled): Engine(cancellation).execute([image], "to:jpg")
    assert not list(image.parent.glob("*.jpg"))


def test_cancel_running_process_stops_its_child(tmp_path):
    import ctypes
    import sys
    import time
    if os.name != "nt": pytest.skip("Contrôle des descendants Windows")
    child_pid = tmp_path / "child.pid"
    code = "import subprocess, sys, pathlib, time; child=subprocess.Popen([sys.executable,'-c','import time; time.sleep(60)']); pathlib.Path(sys.argv[1]).write_text(str(child.pid)); time.sleep(60)"
    engine = Engine()
    timer = threading.Timer(1, engine.cancel.set)
    timer.start()
    start = time.monotonic()
    try:
        with pytest.raises(Cancelled):
            engine.run_process([sys.executable, "-c", code, str(child_pid)])
    finally:
        timer.cancel()
    assert time.monotonic() - start < 12
    assert child_pid.is_file()
    api = ctypes.WinDLL("kernel32", use_last_error=True)
    api.OpenProcess.restype = ctypes.c_void_p
    api.OpenProcess.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
    api.WaitForSingleObject.argtypes = [ctypes.c_void_p, ctypes.c_uint32]
    api.CloseHandle.argtypes = [ctypes.c_void_p]
    handle = api.OpenProcess(0x100000, False, int(child_pid.read_text()))
    if handle:
        try: assert api.WaitForSingleObject(handle, 0) == 0
        finally: api.CloseHandle(handle)


def test_spreadsheet_formula_and_multisheet(tmp_path):
    source = tmp_path / "données.csv"
    source.write_text('Nom;Montant\nTest;42\n=HYPERLINK("https://invalid");0\n', encoding="utf-8")
    out = Engine().execute([source], "to:xlsx")[0]
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.active["A3"].data_type != "f"
    wb.close()
    restored = Engine().execute([out], "to:tsv")[0]
    assert "Test\t42" in Path(restored).read_text(encoding="utf-8-sig")


@pytest.mark.parametrize("ext", ["docx", "html", "epub", "pdf", "txt", "odt", "rtf"])
def test_documents(tmp_path, ext):
    source = tmp_path / "note.md"
    source.write_text("# Zeste\n\nUn document **local**, avec des accents : été.\n", encoding="utf-8")
    outputs = Engine().execute([source], "to:"+ext)
    assert len(outputs) == 1
    result = Path(outputs[0])
    assert result.suffix == "." + ext
    assert not result.name.startswith(".~lock")
    assert result.stat().st_size > 0
    if ext in {"docx", "odt", "epub"}: assert zipfile.is_zipfile(result)
    if ext == "pdf":
        with fitz.open(result) as document: assert "Zeste" in document[0].get_text()


def test_subtitles(tmp_path):
    source = tmp_path / "captions.srt"
    source.write_text("1\n00:00:00,500 --> 00:00:01,500\nBonjour Zeste\n", encoding="utf-8")
    vtt = Path(Engine().execute([source], "to:vtt")[0])
    assert vtt.read_text().startswith("WEBVTT")
    assert "00:00:00.500" in vtt.read_text()
    srt = Path(Engine().execute([vtt], "to:srt")[0])
    assert "Bonjour Zeste" in srt.read_text()


@pytest.mark.skipif(not office_path(), reason="Moteur LibreOffice absent")
def test_office_slides_and_sheets(tmp_path):
    from pptx import Presentation
    from openpyxl import Workbook
    slides = tmp_path / "slides.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Zeste"
    prs.save(slides)
    result = Engine().execute([slides], "to:pdf")[0]
    with fitz.open(result) as doc: assert "Zeste" in doc[0].get_text()
    sheet = tmp_path / "sheet.xlsx"
    wb = Workbook(); wb.active.append(["Test", 42]); wb.save(sheet)
    result = Engine().execute([sheet], "to:pdf")[0]
    with fitz.open(result) as doc: assert "42" in doc[0].get_text()


def test_batch_continues_after_missing_and_corrupt_file(image, tmp_path):
    corrupt = tmp_path / "corrupt.png"
    corrupt.write_text("not an image", encoding="utf8")
    before = hashlib.sha256(image.read_bytes()).digest()
    engine = Engine()
    outputs = engine.execute([tmp_path / "missing.png", image, corrupt, image], "to:jpg")
    assert len(outputs) == 2 and len(set(outputs)) == 2
    assert len(engine.warnings) == 2
    assert "missing.png" in engine.warnings[0] and "corrupt.png" in engine.warnings[1]
    for output in outputs:
        with Image.open(output) as converted: assert converted.size == (320, 240)
    assert hashlib.sha256(image.read_bytes()).digest() == before
    assert not list(tmp_path.glob(".zeste-*"))


def test_batch_all_failed_remains_failure(tmp_path):
    engine = Engine()
    with pytest.raises(ConversionError, match="Aucun fichier du lot"):
        engine.execute([tmp_path / "missing1.png", tmp_path / "missing2.png"], "to:jpg")
    assert len(engine.warnings) == 2


def test_fractional_audio_split_keeps_full_timeline(media, tmp_path):
    outputs = Engine(output_dir=tmp_path).execute([media[1]], "split_media", {"segment": .75})
    durations = [Engine().duration(output) for output in outputs]
    assert len(outputs) == 3
    assert durations == pytest.approx([.75, .75, .5], abs=.02)
    assert sum(durations) == pytest.approx(2, abs=.02)


def test_join_silent_then_audible_video_has_one_complete_audio_track(media, tmp_path):
    import wave
    import numpy as np
    engine = Engine(output_dir=tmp_path)
    silent = tmp_path / "silent.mp4"
    engine.ff(["-i", media[0], "-c:v", "copy", "-an", silent])
    result = Path(engine.execute([silent, media[0]], "join_media")[0])
    info = engine.media_info(result)
    assert info["video_streams"] == 1 and info["audio_streams"] == 1
    assert info["duration"] == pytest.approx(4, abs=.15)
    waveform = tmp_path / "check.wav"
    engine.ff(["-i", result, "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", waveform])
    with wave.open(str(waveform), "rb") as audio:
        samples = np.frombuffer(audio.readframes(audio.getnframes()), dtype=np.int16).astype(float)
    assert np.sqrt(np.mean(samples[8000:16000]**2)) < 5
    assert np.sqrt(np.mean(samples[48000:56000]**2)) > 500


@pytest.mark.parametrize("ext", ["mp4", "webm", "avi", "ogg"])
def test_target_size_uses_bitrate_without_quality_mode(media, tmp_path, ext, monkeypatch):
    engine = Engine(output_dir=tmp_path)
    source = media[1] if ext == "ogg" else media[0]
    if source.suffix != "."+ext:
        source = Path(engine.execute([source], "to:"+ext)[0])
    recorded = []
    run_ff = engine.ff
    def record(args):
        recorded.append(args)
        return run_ff(args)
    monkeypatch.setattr(engine, "ff", record)
    result = Path(engine.execute([source], "compress", {"target_mb": .12})[0])
    assert result.stat().st_size and engine.duration(result) >= 1.8
    assert all(option not in recorded[-1] for option in {"-crf", "-q:v", "-q:a"})
    assert recorded[-1].count("-b:v" if ext != "ogg" else "-b:a") == 1
