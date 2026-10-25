import os
import time
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from pathlib import Path
import pytest
from PIL import Image
from PySide6.QtCore import Qt, QSettings, QMimeData, QUrl, QPoint, QPointF, QCoreApplication, QEvent
from PySide6.QtGui import QDragEnterEvent, QDropEvent, QFontDatabase, QMouseEvent
from PySide6.QtWidgets import QApplication
from PySide6.QtTest import QTest
from zeste.ui import MainWindow
from zeste.ui import OptionsDialog
from zeste.radial import Wheel
from zeste.catalog import operations
from zeste.image_editor import ImageEditor


@pytest.fixture(scope="module")
def app(tmp_path_factory):
    app=QApplication.instance() or QApplication([])
    app.setQuitOnLastWindowClosed(False)
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat,QSettings.Scope.UserScope,str(tmp_path_factory.mktemp("settings")))
    os.environ["ZESTE_CONFIG_DIR"] = str(tmp_path_factory.mktemp("zeste-config"))
    for name in ["segoeui.ttf","segoeuib.ttf","seguisym.ttf"]:
        p=Path("C:/Windows/Fonts")/name
        if p.exists(): QFontDatabase.addApplicationFont(str(p))
    yield app
    for widget in app.topLevelWidgets():
        widget.hide()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()


@pytest.fixture
def photo(tmp_path):
    p=tmp_path/"Exemple été.png"
    Image.new("RGB",(300,200),"coral").save(p)
    return p


def test_drop_selects_sector_and_preserves_source(app,photo):
    original=photo.read_bytes()
    wheel=Wheel();wheel.resize(340,340);wheel.show();QTest.qWait(40)
    got=[];wheel.activated.connect(lambda op,paths:got.append((op,paths)))
    mime=QMimeData();mime.setUrls([QUrl.fromLocalFile(str(photo))])
    enter=QDragEnterEvent(QPoint(170,58),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(wheel,enter);QTest.qWait(20)
    assert enter.isAccepted() and [Path(p) for p in wheel.paths]==[photo]
    op=wheel.visible_items()[0].id
    drop=QDropEvent(QPointF(170,58),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(wheel,drop)
    assert drop.isAccepted() and got[0][0]==op and [Path(p) for p in got[0][1]]==[photo]
    assert photo.read_bytes()==original
    wheel.close()


def test_wheel_keyboard_and_pages(app,photo):
    wheel=Wheel(popup=True);wheel.reduce_motion=True;wheel.set_files([str(photo)]);wheel.show();QTest.qWait(20)
    QTest.keyClick(wheel,Qt.Key.Key_Tab)
    assert wheel.mode=="tools"
    assert wheel.visible_items()[-1].id=="next"
    wheel.next_page();assert wheel.page==1
    QTest.keyClick(wheel,Qt.Key.Key_Escape);assert not wheel.isVisible()


def test_actions_empty_selection_and_cardinality(photo):
    assert operations([],"actions")==[]
    assert not {"join_media","merge_pdf","collage"}&{x.id for x in operations([str(photo)],"tools",True)}
    assert 'edit_metadata' in {x.id for x in operations([str(photo)],'tools',True)}
    assert "collage" in {x.id for x in operations([str(photo),str(photo)],"tools",True)}


@pytest.mark.parametrize("label,page", [("Réglages",2),("Fonctionnalités…",3),("Atelier avancé",0),("Derniers résultats",1)])
def test_tray_action_opens_requested_page(app,label,page):
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.navigate(3)
    action=next(a for a in window.tray.contextMenu().actions() if a.text()==label)
    action.trigger();app.processEvents()
    assert window.isVisible() and window.stack.currentIndex()==page
    window.hide()
    action.trigger();app.processEvents()
    assert window.isVisible() and window.stack.currentIndex()==page
    app.removeEventFilter(window)
    window.deleteLater()
    QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete)


def test_wheel_center_drag_moves_without_changing_mode_or_running_action(app,photo):
    wheel=Wheel(popup=True);wheel.reduce_motion=True;wheel.move(40,40)
    wheel.set_files([str(photo)]);wheel.show();QTest.qWait(30)
    fired=[];wheel.activated.connect(lambda *args:fired.append(args))
    center=wheel.rect().center();origin=wheel.pos();global_start=wheel.mapToGlobal(center)
    QTest.mousePress(wheel,Qt.MouseButton.LeftButton,pos=center)
    delta=QPoint(95,65)
    move=QMouseEvent(QEvent.Type.MouseMove,QPointF(center+delta),QPointF(global_start+delta),
                     Qt.MouseButton.NoButton,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(wheel,move)
    QTest.mouseRelease(wheel,Qt.MouseButton.LeftButton,pos=center)
    assert wheel.pos()==origin+delta
    assert wheel.mode=="quick" and not fired and wheel.isVisible()
    assert wheel.move_origin is None and not wheel.moving
    # A simple centre click still switches menus after a move.
    QTest.mouseClick(wheel,Qt.MouseButton.LeftButton,pos=center)
    assert wheel.mode=="tools"
    wheel.close()


def test_wheel_drag_stays_on_screen_and_hide_cancels_move(app):
    wheel=Wheel(popup=True);wheel.reduce_motion=True;wheel.show();QTest.qWait(20)
    center=wheel.rect().center();global_start=wheel.mapToGlobal(center)
    QTest.mousePress(wheel,Qt.MouseButton.LeftButton,pos=center)
    move=QMouseEvent(QEvent.Type.MouseMove,QPointF(center+QPoint(50000,50000)),QPointF(global_start+QPoint(50000,50000)),
                     Qt.MouseButton.NoButton,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(wheel,move)
    assert wheel.screen().availableGeometry().contains(wheel.frameGeometry())
    wheel.hide()
    assert wheel.move_origin is None and not wheel.moving
    wheel.close()


def test_result_gallery_tracks_real_outputs_and_respects_feature_switches(app,photo,tmp_path):
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue("global_drag",False)
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.add_files([str(photo)])
    assert window.quick_edit.isEnabled()
    window.settings.setValue("features/image/edit_image",False);window.refresh_features()
    assert not window.quick_edit.isEnabled()
    missing=tmp_path/"removed.png"
    window.history_records=[{"time":"Test","operation":"to:png","outputs":[str(photo),str(photo),str(missing)],"success":True,"detail":""}]
    window.render_history()
    assert window.gallery.count()==1
    item=window.gallery.item(0)
    assert item.data(Qt.ItemDataRole.UserRole)==str(photo) and not item.icon().isNull()
    window.active_job=([str(photo)],"to:png")
    window.on_finished([str(photo)],[])
    assert window.wheel.progress_value==100 and not window.wheel.busy
    assert window.gallery.count()==1
    window.settings.setValue("features/image/edit_image",True)
    app.removeEventFilter(window);window.deleteLater()
    QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete)


def test_image_editor_coordinates_feed_saved_copy(app,photo,tmp_path):
    editor=ImageEditor(str(photo));editor.show();QTest.qWait(30)
    canvas=editor.canvas;canvas.tool="redact";rect=canvas.image_rect()
    a=QPoint(round(rect.left()+rect.width()*.1),round(rect.top()+rect.height()*.1))
    b=QPoint(round(rect.left()+rect.width()*.4),round(rect.top()+rect.height()*.4))
    QTest.mousePress(canvas,Qt.MouseButton.LeftButton,pos=a);QTest.mouseMove(canvas,b);QTest.mouseRelease(canvas,Qt.MouseButton.LeftButton,pos=b)
    from zeste.engine import Engine
    out=Engine(output_dir=tmp_path).execute([photo],"edit_image",editor.options())[0]
    with Image.open(out) as im: assert im.getpixel((60,40))[:3]==(0,0,0)
    with Image.open(photo) as im: assert im.getpixel((60,40))[:3]!=(0,0,0)
    editor.close()


def test_window_queue_conversion_history_and_error(app,photo,tmp_path):
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue("global_drag",False)
    window=MainWindow();window.watcher.timer.stop();window.tray.hide();window.show()
    window.settings.setValue("output",str(tmp_path/"results"))
    window.add_files([str(photo)])
    assert window.run_button.isEnabled()
    window.show_popup()
    QTest.keyClick(window,Qt.Key.Key_Tab)
    assert window.popup.mode=="tools"
    QTest.keyClick(window.popup,Qt.Key.Key_Tab)
    assert window.popup.mode=="actions"
    QTest.keyClick(window,Qt.Key.Key_Escape)
    assert not window.popup.isVisible()
    window.pending.extend([([str(photo)],"to:webp",{}),([str(tmp_path/"missing.png")],"to:jpg",{})])
    window.start_next()
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        QTest.qWait(25)
        time.sleep(.01)  # Let the Python worker acquire the GIL between Qt waits.
        if window.thread is None and not window.pending: break
    assert window.thread is None
    assert len(list((tmp_path/"results").glob("*.webp")))==1
    assert len(window.history_records)==2
    assert window.history_records[0]["success"] is False
    assert window.history_records[1]["success"] is True
    window.set_appearance("light",True)
    assert window.wheel.light and window.popup.light
    window.hide();window.popup.close();window.result_panel.close();window.deleteLater()
    QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete)
    app.processEvents()


def test_dialog_options_really_apply_to_engine(app,photo,tmp_path):
    from zeste.engine import Engine
    import pymupdf as fitz
    dialog = OptionsDialog("frame_image",[str(photo)],None)
    framed=Engine().execute([photo],"frame_image",dialog.values())[0]
    with Image.open(framed) as im: assert im.size==(460,360)
    dialog.close()
    jpeg=tmp_path/"photo.jpg"
    Image.open(photo).save(jpeg)
    dialog=OptionsDialog("edit_metadata",[str(jpeg)],None)
    dialog.fields["author"].setText("Auteur UI")
    out=Engine().execute([jpeg],"edit_metadata",dialog.values())[0]
    with Image.open(out) as im: assert im.getexif()[315]=="Auteur UI"
    dialog.close()
    pdf=tmp_path/"source.pdf"
    with fitz.open() as doc: doc.new_page();doc.save(pdf)
    dialog=OptionsDialog("edit_metadata",[str(pdf)],None)
    dialog.fields["subject"].setText("Sujet UI")
    out=Engine().execute([pdf],"edit_metadata",dialog.values())[0]
    with fitz.open(out) as doc: assert doc.metadata["subject"]=="Sujet UI"
    dialog.close()


@pytest.mark.parametrize("operation",["compress","to:jpg"])
def test_hidden_app_drop_runs_without_opening_workshop(app,photo,tmp_path,operation):
    import math
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue("hud",False);settings.setValue("global_drag",False)
    original=photo.read_bytes()
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.settings.setValue("output",str(tmp_path/"converted"))
    wheel=window.popup;wheel.reduce_motion=True;wheel.show()
    mime=QMimeData();mime.setUrls([QUrl.fromLocalFile(str(photo))])
    enter=QDragEnterEvent(QPoint(160,160),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.ControlModifier)
    QApplication.sendEvent(wheel,enter);QTest.qWait(30)
    index=next(i for i,op in enumerate(wheel.visible_items()) if op.id==operation)
    angle=math.pi/2-math.tau*index/len(wheel.visible_items())
    point=QPointF(160+math.cos(angle)*105.5,160-math.sin(angle)*105.5)
    drop=QDropEvent(point,Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(wheel,drop)
    assert drop.isAccepted()
    deadline=time.monotonic()+30
    while time.monotonic()<deadline:
        QTest.qWait(20);time.sleep(.01)
        if window.history_records and window.thread is None: break
    assert window.history_records[0]["success"]
    result=Path(window.history_records[0]["outputs"][0])
    with Image.open(result) as converted: assert converted.size==(300,200)
    assert photo.read_bytes()==original
    assert not window.isVisible() and not wheel.isVisible() and not window.result_panel.isVisible()
    window.popup.close();window.result_panel.close();window.deleteLater()
    QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_glass_wheel_has_transparent_corners_and_reference_teal(app,photo):
    wheel=Wheel(popup=True);wheel.reduce_motion=True;wheel.set_files([str(photo)]);wheel.show()
    wheel.hovered=0;wheel.update();QTest.qWait(30)
    rendered=wheel.grab().toImage()
    assert rendered.pixelColor(0,0).alpha()==0
    assert 0<rendered.pixelColor(265,160).alpha()<255
    assert rendered.pixelColor(160,30).name()=="#a4e7e3"
    wheel.close()


def drop_quick(window, path, operation):
    import math
    wheel=window.popup;wheel.reduce_motion=True;wheel.set_mode('quick');wheel.show()
    mime=QMimeData();mime.setUrls([QUrl.fromLocalFile(str(path))])
    enter=QDragEnterEvent(QPoint(160,160),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.ControlModifier)
    QApplication.sendEvent(wheel,enter);QTest.qWait(25)
    index=next(i for i,op in enumerate(wheel.visible_items()) if op.id==operation)
    angle=math.pi/2-math.tau*index/len(wheel.visible_items())
    point=QPointF(160+math.cos(angle)*105,160-math.sin(angle)*105)
    drop=QDropEvent(point,Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(wheel,drop)
    assert drop.isAccepted()


def await_job(window, count):
    deadline=time.monotonic()+30
    while time.monotonic()<deadline:
        QTest.qWait(20);time.sleep(.01)
        if len(window.history_records)>=count and window.thread is None:return
    pytest.fail('The wheel job did not finish or the queue stayed blocked')


@pytest.mark.parametrize('ext',['pptx','png','jpg'])
def test_pdf_quick_drop_converts_without_modal_dialog(app,tmp_path,monkeypatch,ext):
    import pymupdf as fitz
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue('global_drag',False)
    source=tmp_path/'pages.pdf'
    with fitz.open() as doc:
        doc.new_page().insert_text((20,30),'First page')
        doc.new_page().insert_text((20,30),'Second page')
        doc.save(source)
    before=source.read_bytes()
    def unexpected_dialog(*args,**kwargs): raise AssertionError('A quick conversion must not open a modal dialog')
    monkeypatch.setattr('zeste.ui.OptionsDialog',unexpected_dialog)
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.settings.setValue('output',str(tmp_path/'out'))
    drop_quick(window,source,'to:'+ext);await_job(window,1)
    record=window.history_records[0]
    assert record['success'],record['detail']
    if ext=='pptx':
        from pptx import Presentation
        slides=Presentation(record['outputs'][0]).slides
        assert len(slides)==2 and all(len(slide.shapes)==1 for slide in slides)
    else:assert len(record['outputs'])==2
    assert QApplication.activeModalWidget() is None
    assert not window.isVisible() and not window.popup.isVisible() and not window.result_panel.isVisible()
    assert source.read_bytes()==before
    window.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_wheel_accepts_next_file_after_failed_powerpoint(app,tmp_path,photo):
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue('global_drag',False)
    broken=tmp_path/'broken.pdf';broken.write_bytes(b'Not a PDF')
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.settings.setValue('output',str(tmp_path/'out'))
    drop_quick(window,broken,'to:pptx');await_job(window,1)
    assert not window.history_records[0]['success']
    assert 'Échec' in window.tray.toolTip()
    drop_quick(window,photo,'to:webp');await_job(window,2)
    assert window.history_records[0]['success']
    assert len(list((tmp_path/'out').glob('*.webp')))==1
    assert window.thread is None and not window.pending and QApplication.activeModalWidget() is None
    window.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_preparation_error_is_reported_and_does_not_block_next_drop(app,tmp_path,photo,monkeypatch):
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue('global_drag',False)
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.settings.setValue('output',str(tmp_path/'out'))
    def fail(*args,**kwargs):raise RuntimeError('Preparation failed')
    with monkeypatch.context() as scoped:
        scoped.setattr(window,'request_operation',fail)
        drop_quick(window,photo,'to:webp');await_job(window,1)
    assert not window.history_records[0]['success']
    drop_quick(window,photo,'to:webp');await_job(window,2)
    assert window.history_records[0]['success']
    window.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_advanced_modal_never_leaves_a_disabled_wheel_above_it(app):
    from PySide6.QtWidgets import QDialog
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    dialog=QDialog(window);dialog.setModal(True);dialog.show();QTest.qWait(20)
    window.popup.show();window.watcher.poll()
    assert not window.popup.isVisible()
    dialog.reject()
    assert QApplication.activeModalWidget() is None
    window.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_feature_switches_persist_filter_both_wheels_and_block_dispatch(app,photo):
    from zeste.preferences import load_settings
    settings=load_settings();settings.clear();settings.setValue('global_drag',False)
    window=MainWindow();window.watcher.timer.stop();window.tray.hide()
    window.add_files([str(photo)]);window.wheel.set_mode('tools');window.popup.set_files([str(photo)]);window.popup.set_mode('tools')
    window.features.items['image/crop_image'].setCheckState(0,Qt.CheckState.Unchecked)
    assert 'crop_image' not in {op.id for op in window.wheel.items}
    assert 'crop_image' not in {op.id for op in window.popup.items}
    assert load_settings().value('features/image/crop_image',True,bool) is False
    window.request_operation('crop_image',[str(photo)])
    assert not window.pending and window.thread is None and QApplication.activeModalWidget() is None
    window.features.items['image/crop_image'].setCheckState(0,Qt.CheckState.Checked)
    assert 'crop_image' in {op.id for op in window.popup.items}
    window.features.search.setText('audio')
    assert not window.features.items['audio/trim'].isHidden() and window.features.items['image/crop_image'].isHidden()
    window.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_center_hover_keeps_destinations_until_explicit_tab(app,photo):
    wheel=Wheel(popup=True);wheel.reduce_motion=True;wheel.set_files([str(photo)]);wheel.show();QTest.qWait(20)
    before=[op.id for op in wheel.visible_items()]
    wheel.hover(QPointF(160,160),True);QTest.qWait(1000)
    assert wheel.mode=='quick' and [op.id for op in wheel.visible_items()]==before
    QTest.keyClick(wheel,Qt.Key.Key_Tab)
    assert wheel.mode=='tools'
    wheel.close()


def test_image_modes_numeric_crop_and_disabled_annotations(app,photo,tmp_path):
    from zeste.engine import Engine
    editor=ImageEditor(str(photo),initial='crop_image',enabled=lambda key:key!='image/annotate_image')
    assert editor.sections.currentData()=='crop_image' and editor.sections.findData('annotate_image')==-1
    editor.crop_values['x'].setValue(20);editor.crop_values['y'].setValue(10)
    editor.crop_values['width'].setValue(100);editor.crop_values['height'].setValue(70)
    result=Engine(output_dir=tmp_path/'out').execute([photo],'crop_image',editor.options())[0]
    with Image.open(result) as im:assert im.size==(100,70)
    editor.close()


def test_pdf_organizer_visual_changes_export_in_selected_order(app,tmp_path):
    import pymupdf as fitz
    from zeste.pdf_organizer import PdfOrganizer
    from zeste.engine import Engine
    path=tmp_path/'pages.pdf'
    with fitz.open() as doc:
        for text in ['First','Second']:doc.new_page().insert_text((20,30),text)
        doc.save(path)
    dialog=PdfOrganizer(path);dialog.pages.item(1).setSelected(True);dialog.rotate();dialog.duplicate()
    dialog.pages.insertItem(0,dialog.pages.takeItem(2))
    result=Engine(output_dir=tmp_path/'out').execute([path],'reorder_pdf',dialog.options())[0]
    with fitz.open(result) as doc:
        assert [p.get_text().strip() for p in doc]==['Second','First','Second']
        assert [p.rotation for p in doc]==[90,0,90]
    dialog.close()

@pytest.mark.parametrize('light', [False, True])
def test_style_switch_roundtrip_persists_and_preserves_work(app, photo, tmp_path, monkeypatch, light):
    from zeste.preferences import load_settings
    monkeypatch.setenv('ZESTE_CONFIG_DIR', str(tmp_path/'theme-profile'))
    settings=load_settings(); settings.setValue('light',light); settings.setValue('glass_strength',80)
    settings.setValue('features/image/crop_image',False); settings.sync()
    w=MainWindow();w.watcher.timer.stop();w.tray.hide();w.add_files([str(photo)])
    w.show_page(2);app.processEvents()
    original_style=app.styleSheet(); original_wheel=w.wheel.grab().toImage()
    original_icon=w.tray.icon().pixmap(32,32).toImage()
    QTest.mouseClick(w.style_switch, Qt.MouseButton.LeftButton);app.processEvents()
    assert w.style_switch.isChecked() and w.popup.color_style==w.wheel.color_style=='melina'
    assert load_settings().value('color_style')=='melina'
    assert w.centralWidget().color_style==w.result_panel.color_style==w.result_panel.mark.color_style=='melina'
    assert not w.theme_picker.isEnabled() and not w.feature_enabled('image/crop_image')
    assert w.selected_paths()==[str(photo)] and w.popup.glass_strength==.8
    assert w.wheel.grab().toImage()!=original_wheel and w.tray.icon().pixmap(32,32).toImage()!=original_icon
    assert w.windowTitle().startswith('Meloo') and '#4a302d' in app.styleSheet()
    w2=MainWindow();w2.watcher.timer.stop();w2.tray.hide()
    assert w2.style_switch.isChecked() and w2.popup.color_style=='melina'
    w2.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete)
    w.style_switch.setFocus();QTest.keyClick(w.style_switch,Qt.Key.Key_Space);app.processEvents()
    assert not w.style_switch.isChecked() and load_settings().value('color_style')=='orange'
    assert w.theme_picker.isEnabled() and w.settings.value('light',False,bool)==light
    assert app.styleSheet()==original_style and w.wheel.grab().toImage()==original_wheel
    assert w.tray.icon().pixmap(32,32).toImage()==original_icon
    assert not w.feature_enabled('image/crop_image') and w.selected_paths()==[str(photo)]
    w.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_melina_quick_drop_converts_and_keeps_transparent_corners(app,photo,tmp_path,monkeypatch):
    from zeste.preferences import load_settings
    monkeypatch.setenv('ZESTE_CONFIG_DIR',str(tmp_path/'profile'))
    settings=load_settings();settings.setValue('color_style','melina');settings.setValue('output',str(tmp_path/'out'));settings.sync()
    w=MainWindow();w.watcher.timer.stop();w.tray.hide();w.popup.reduce_motion=True
    w.popup.show();app.processEvents()
    assert w.popup.grab().toImage().pixelColor(0,0).alpha()==0
    drop_quick(w,photo,'to:webp');await_job(w,1)
    assert w.history_records[0]['success'] and w.popup.color_style=='melina'
    assert len(list((tmp_path/'out').glob('*.webp')))==1
    w.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()


def test_destination_slots_stay_fixed_across_input_formats_and_feature_changes(app):
    wheel=Wheel(popup=True);wheel.reduce_motion=True
    wheel.set_files(['sample.png']);png=wheel.visible_items()
    wheel.set_files(['sample.jpg']);jpg=wheel.visible_items()
    assert len(png)==len(jpg)==8 and png[0].id==jpg[0].id=='create_zip'
    assert png[1].id==jpg[1].id=='compress'
    assert png[2].id=='to:jpg' and jpg[2].id==''
    assert png[3].id=='' and jpg[3].id=='to:png'
    assert png[4].id==jpg[4].id=='to:webp'
    wheel.feature_enabled=lambda key:key!='convert/webp'
    wheel.refresh();disabled=wheel.visible_items()
    assert disabled[4].id=='' and disabled[5].id=='to:pdf'
    wheel.next_page()
    assert wheel.visible_items()[0].id=='create_zip' and len(wheel.visible_items())==8
    wheel.close()


def test_next_page_hover_advances_once_until_pointer_leaves(app,photo):
    import math
    wheel=Wheel(popup=True);wheel.reduce_motion=True;wheel.set_files([str(photo)]);wheel.show();QTest.qWait(30)
    angle=math.pi/2-math.tau*7/8
    point=QPointF(160+math.cos(angle)*105,160-math.sin(angle)*105)
    wheel.hover(point,True);QTest.qWait(1000)
    assert wheel.page==1
    for _ in range(5): wheel.hover(point,True);QTest.qWait(220)
    assert wheel.page==1
    wheel.hover(QPointF(160,160),True);wheel.hover(point,True);QTest.qWait(1000)
    assert wheel.page==0
    wheel.close()


def test_create_zip_drop_groups_original_file_without_opening_dialog(app,photo,tmp_path,monkeypatch):
    import zipfile
    from zeste.preferences import load_settings
    monkeypatch.setenv('ZESTE_CONFIG_DIR',str(tmp_path/'zip-profile'))
    settings=load_settings();settings.setValue('output',str(tmp_path/'out'));settings.sync()
    w=MainWindow();w.watcher.timer.stop();w.tray.hide()
    drop_quick(w,photo,'create_zip');await_job(w,1)
    assert w.history_records[0]['success']
    with zipfile.ZipFile(w.history_records[0]['outputs'][0]) as z:
        assert z.namelist()==[photo.name] and z.read(photo.name)==photo.read_bytes()
    assert not w.isVisible() and QApplication.activeModalWidget() is None
    w.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents()
