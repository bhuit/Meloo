"""Exercise the actual windowed portable, with isolated profiles and no venv PATH."""
import array
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import tempfile
import time
import wave

import pytest
from PIL import Image
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[1]
EXE = Path(os.environ.get("MELOO_TEST_EXE", ROOT / "dist/Zeste/Meloo.exe"))
pytestmark = pytest.mark.skipif(not EXE.is_file(), reason="Portable not built")


def test_packaged_photo_editor_pipeline_flattens_redaction(tmp_path):
    source=tmp_path/'photo.png'; Image.new('RGB',(150,100),'coral').save(source)
    result=convert(tmp_path,source,'redact_image',{'exposure':.4,'crop':[0,0,100,80],'annotations':[{'type':'redact','box':[10,10,40,40]}]})[0]
    with Image.open(result) as image:
        assert image.size==(100,80) and image.getpixel((20,20))[:3]==(0,0,0)


def test_packaged_visual_pdf_organization(tmp_path):
    source=tmp_path/'pages.pdf'
    with fitz.open() as doc:
        for text in ['FIRST','SECOND']:doc.new_page().insert_text((20,30),text)
        doc.save(source)
    result=convert(tmp_path,source,'reorder_pdf',{'page_plan':[{'page':1,'rotation':90},{'page':0,'rotation':0}]})[0]
    with fitz.open(result) as doc:
        assert [p.get_text().strip() for p in doc]==['SECOND','FIRST']
        assert [p.rotation for p in doc]==[90,0]


def test_packaged_office_metadata_preserves_document(tmp_path):
    from docx import Document
    source=tmp_path/'document.docx'; doc=Document();doc.add_paragraph('Contenu original');doc.save(source)
    result=convert(tmp_path,source,'edit_metadata',{'title':'Titre modifié','author':'Auteur'})[0]
    output=Document(result)
    assert output.core_properties.title=='Titre modifié' and output.paragraphs[0].text=='Contenu original'


def run_portable(tmp_path, *args):
    profile = tmp_path / "profile"
    env = dict(os.environ)
    for name in ["IMAGEIO_FFMPEG_EXE", "PYPANDOC_PANDOC", "ZESTE_LIBREOFFICE", "PYTHONPATH", "PYTHONHOME"]:
        env.pop(name, None)
    # Native LibreOffice uses deep profile subdirectories; keep the isolated
    # scratch root short, as with the normal Windows per-user TEMP directory.
    scratch_root = ROOT / "build/tmp"
    scratch_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="p-", dir=scratch_root) as scratch:
        env.update(QT_QPA_PLATFORM="offscreen", ZESTE_CONFIG_DIR=str(profile), LOCALAPPDATA=str(tmp_path / "local"), APPDATA=str(tmp_path / "roaming"), TEMP=scratch, TMP=scratch, PATH=str(Path(os.environ.get("WINDIR", "C:/Windows")) / "System32"))
        result = subprocess.run([str(EXE), *map(str, args)], cwd=tmp_path, env=env, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=240, creationflags=0x08000000)
    error_path = tmp_path / "local/Zeste/startup-error.log"
    error = error_path.read_text(encoding="utf8") if error_path.exists() else result.stderr.decode(errors="replace")
    assert result.returncode == 0, error or f"Portable exited {result.returncode}"
    assert not error_path.exists(), error
    return result


def convert(tmp_path, source, operation, options=None):
    output = tmp_path / "result"
    output.mkdir()
    args = ["--convert", operation, "--output", output]
    if options is not None: args += ["--options", json.dumps(options)]
    run_portable(tmp_path, *args, source)
    files = list(output.iterdir())
    assert files and all(path.stat().st_size for path in files)
    return files


def test_packaged_smoke_has_real_zero_exit(tmp_path):
    run_portable(tmp_path, "--smoke")


@pytest.mark.parametrize("color_style", ["orange", "melina"])
def test_packaged_background_has_tray_and_no_main_window(tmp_path, color_style):
    profile=tmp_path/"resident-profile";profile.mkdir()
    (profile/"settings.ini").write_text(f"[General]\nbackground_hint_shown=true\nhud=false\ncolor_style={color_style}\n",encoding="utf8")
    env=dict(os.environ)
    env.update(QT_QPA_PLATFORM="windows",ZESTE_CONFIG_DIR=str(profile),
               LOCALAPPDATA=str(tmp_path/"local"),APPDATA=str(tmp_path/"roaming"),
               PATH=str(Path(os.environ.get("WINDIR","C:/Windows"))/"System32"))
    startup=subprocess.STARTUPINFO()
    startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    startup.wShowWindow=0  # Same hidden launch as the installed updater.
    process=subprocess.Popen([str(EXE)],cwd=tmp_path,env=env,creationflags=0x08000000,startupinfo=startup)
    status_path=tmp_path/"resident-status.json"
    try:
        status={}
        for attempt in range(4):
            time.sleep(.5)
            subprocess.run([str(EXE),"--status-file",str(status_path)],cwd=tmp_path,env=env,timeout=45,check=True,creationflags=0x08000000)
            status=json.loads(status_path.read_text(encoding="utf8"))
            if status.get("running"): break
        assert process.poll() is None
        assert status["version"]=="0.4.2"
        assert status["app_name"]=="Meloo" and status["color_style"]==color_style
        assert status["tray_visible"] and status["watcher_active"]
        assert status["window_visible"] is False
        assert status["trigger"]=="Ctrl" and status["default_mode"]=="quick"
        assert status["input_backend"]=="raw-input" and status["input_error"]==""
        assert status["window_native_visible"] is False
        subprocess.run([str(EXE.with_name("Zeste.exe")),"--window"],cwd=tmp_path,env=env,timeout=45,check=True,creationflags=0x08000000)
        time.sleep(.3)
        subprocess.run([str(EXE),"--status-file",str(status_path)],cwd=tmp_path,env=env,timeout=45,check=True,creationflags=0x08000000)
        status=json.loads(status_path.read_text(encoding="utf8"))
        assert status["window_visible"] and status["window_native_visible"], status
        assert status["platform"]=="windows" and status["current_page"]==0
        subprocess.run([str(EXE),"--quit"],cwd=tmp_path,env=env,timeout=45,check=True,creationflags=0x08000000)
        assert process.wait(timeout=15)==0
    finally:
        if process.poll() is None:
            process.kill();process.wait(timeout=10)


def test_packaged_inventory_and_source_identity():
    manifest = json.loads((EXE.parent / ".zeste-generated.json").read_text(encoding="utf8"))
    assert manifest["version"] == "0.4.2"
    assert hashlib.sha256(EXE.read_bytes()).digest()==hashlib.sha256(EXE.with_name("Zeste.exe").read_bytes()).digest()
    for relative, expected in manifest["source_sha256"].items():
        assert hashlib.sha256((ROOT / relative).read_bytes()).hexdigest() == expected
    names = [path.name.lower() for path in (EXE.parent / "_internal").rglob("*")]
    assert not any(name.startswith(("qt6qml", "qt6quick", "qt6webengine", "qt63d", "onnxruntime", "pymupdf4llm", "pymupdf_layout")) for name in names)
    assert "icuuc.dll" not in names  # Qt uses the Windows ICU API, not Poppler's.
    assert not list((EXE.parent / "tools").rglob("*.msi"))
    assert (EXE.parent / "tools/share/Zeste.Share.exe").is_file()


@pytest.mark.parametrize("ext", ["jpg", "heic"])
def test_packaged_image_codecs(tmp_path, ext):
    source = tmp_path / "image été.png"
    Image.new("RGB", (96, 64), "#ff5c00").save(source)
    output = convert(tmp_path, source, "to:" + ext)[0]
    import pillow_heif
    pillow_heif.register_heif_opener()
    with Image.open(output) as image: assert image.size == (96, 64)


def test_packaged_docx_to_pdf_uses_private_office(tmp_path):
    from docx import Document
    source = tmp_path / "document.docx"
    document = Document()
    document.add_heading("Zeste portable", 0)
    document.add_paragraph("Document local avec accents : été.")
    document.save(source)
    output = convert(tmp_path, source, "to:pdf")[0]
    with fitz.open(output) as pdf: assert "Zeste portable" in pdf[0].get_text()


def test_packaged_pdf_to_docx_without_optional_ai_modules(tmp_path):
    from docx import Document
    source = tmp_path / "document.pdf"
    with fitz.open() as document:
        document.new_page().insert_text((50, 60), "Zeste portable test")
        document.save(source)
    output = convert(tmp_path, source, "to:docx")[0]
    assert "Zeste portable test" in " ".join(p.text for p in Document(output).paragraphs)


def test_packaged_markdown_to_epub_uses_private_pandoc(tmp_path):
    import zipfile
    source = tmp_path / "document.md"
    source.write_text("# Zeste portable\n\nUn **livre** local.", encoding="utf8")
    output = convert(tmp_path, source, "to:epub")[0]
    with zipfile.ZipFile(output) as archive: assert archive.read("mimetype") == b"application/epub+zip"


def test_packaged_audio_uses_private_ffmpeg(tmp_path):
    source = tmp_path / "audio.wav"
    with wave.open(str(source), "wb") as audio:
        audio.setparams((1, 2, 16000, 32000, "NONE", "not compressed"))
        audio.writeframes(array.array("h", (round(3000 * math.sin(i * 2 * math.pi * 440 / 16000)) for i in range(32000))).tobytes())
    output = convert(tmp_path, source, "to:mp3")[0]
    assert output.stat().st_size > 10000 and output.suffix == ".mp3"


def test_packaged_7zip_roundtrip(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("Zeste archive", encoding="utf8")
    archive = convert(tmp_path, source, "to:7z")[0]
    restored = tmp_path / "restored"
    run_portable(tmp_path, "--convert", "extract", "--output", restored, archive)
    assert next(restored.rglob("source.txt")).read_bytes() == source.read_bytes()


def test_packaged_pdf_compression_preserves_content_and_reduces_size(tmp_path):
    source=tmp_path/'document.pdf'
    with fitz.open() as doc:
        page=doc.new_page()
        for i in range(30):page.insert_text((20,30+i*20),f'Compression test {i}')
        doc.save(source)
    before=source.read_bytes()
    output=convert(tmp_path,source,'compress')[0]
    assert output.stat().st_size<len(before)
    with fitz.open(source) as a,fitz.open(output) as b:
        assert a[0].get_text()==b[0].get_text()
        assert a[0].get_pixmap().samples==b[0].get_pixmap().samples
    assert source.read_bytes()==before


def test_packaged_pdf_to_powerpoint_contains_each_page(tmp_path):
    from pptx import Presentation
    source=tmp_path/'pages.pdf'
    with fitz.open() as doc:
        doc.new_page().insert_text((30,30),'First page')
        doc.new_page().insert_text((30,30),'Second page')
        doc.save(source)
    output=convert(tmp_path,source,'to:pptx')[0]
    presentation=Presentation(output)
    assert len(presentation.slides)==2
    assert all(len(slide.shapes)==1 for slide in presentation.slides)


def test_packaged_create_zip_keeps_mixed_originals_intact(tmp_path):
    import zipfile
    source=tmp_path/'image.jpg';source.write_bytes(b'image bytes')
    nested=tmp_path/'existing.zip'
    with zipfile.ZipFile(nested,'w') as archive: archive.writestr('document.txt',b'original')
    output=tmp_path/'out';output.mkdir()
    run_portable(tmp_path,'--convert','create_zip','--output',output,source,nested)
    files=list(output.glob('*.zip'));assert len(files)==1
    with zipfile.ZipFile(files[0]) as archive:
        assert archive.read(source.name)==source.read_bytes()
        assert archive.read(nested.name)==nested.read_bytes()
