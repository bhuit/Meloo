import json
import zipfile
from pathlib import Path
import numpy as np
import pytest
from PIL import Image
import pymupdf as fitz
from zeste.engine import Engine, ConversionError
from zeste.photo import adjust_photo, render_photo, ADJUSTMENTS
from zeste.catalog import operations, feature_key


@pytest.fixture
def photo(tmp_path):
    data=np.random.default_rng(7).integers(20,230,(120,180,4),dtype='uint8'); data[:,:,3]=170
    path=tmp_path/'image.png'; Image.fromarray(data).save(path); return path


@pytest.mark.parametrize('key,value',[(key,(default+high)/2) for key,title,default,low,high,step in ADJUSTMENTS])
def test_photo_adjustment_changes_pixels_preserves_alpha(photo,key,value):
    with Image.open(photo) as original:
        output=adjust_photo(original,{key:value})
        assert output.size==original.size
        assert output.getchannel('A').tobytes()==original.getchannel('A').tobytes()
        assert output.convert('RGB').tobytes()!=original.convert('RGB').tobytes()


def test_annotation_crop_and_redaction_are_flattened(photo,tmp_path):
    original=photo.read_bytes()
    output=Path(Engine(output_dir=tmp_path/'out').execute([photo],'redact_image',{'crop':[10,10,110,110],'annotations':[{'type':'redact','box':[10,10,50,50]},{'type':'arrow','box':[60,60,90,90],'color':'#00ff00'}]})[0])
    with Image.open(output) as im:
        assert im.size==(100,100) and im.getpixel((20,20))==(0,0,0,255)
        assert im.getpixel((89,89))[1]==255
    assert photo.read_bytes()==original


def test_arbitrary_file_compression_and_metadata_preserve_source(tmp_path):
    source=tmp_path/'data.custom'; source.write_bytes(b'Important original content\n'*200)
    original=source.read_bytes(); engine=Engine(output_dir=tmp_path/'out')
    archive=engine.execute([source],'compress')[0]
    with zipfile.ZipFile(archive) as z: assert z.read(source.name)==original
    metadata=engine.execute([source],'edit_metadata',{'title':'Notes'})[0]
    assert json.loads(Path(metadata).read_text(encoding='utf8'))['storage']=='sidecar'
    assert source.read_bytes()==original


def test_png_and_office_metadata_are_embedded_without_content_changes(photo,tmp_path):
    engine=Engine(output_dir=tmp_path/'out')
    output=engine.execute([photo],'edit_metadata',{'title':'Photo','author':'Auteur'})[0]
    with Image.open(photo) as original,Image.open(output) as modified:
        assert original.tobytes()==modified.tobytes() and modified.text['Title']=='Photo'
    from docx import Document
    source=tmp_path/'doc.docx'; doc=Document(); doc.add_paragraph('Le texte reste intact.'); doc.save(source)
    before=source.read_bytes(); output=engine.execute([source],'edit_metadata',{'author':'Auteur','title':'Titre'})[0]
    edited=Document(output); assert edited.core_properties.author=='Auteur' and edited.paragraphs[0].text=='Le texte reste intact.'
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(output) as b:
        assert a.namelist()==b.namelist()
        assert all(a.read(name)==b.read(name) for name in a.namelist() if name!='docProps/core.xml')
    assert source.read_bytes()==before


def test_visual_pdf_plan_duplicates_rotates_reorders_without_altering_source(tmp_path):
    source=tmp_path/'pages.pdf'
    with fitz.open() as doc:
        for text in ['ONE','TWO','THREE']: doc.new_page().insert_text((20,30),text)
        doc.save(source)
    before=source.read_bytes()
    output=Engine(output_dir=tmp_path/'out').execute([source],'reorder_pdf',{'page_plan':[{'page':2,'rotation':90},{'page':0,'rotation':0},{'page':2,'rotation':180}]})[0]
    with fitz.open(output) as doc:
        assert [p.get_text().strip() for p in doc]==['THREE','ONE','THREE']
        assert [p.rotation for p in doc]==[90,0,180]
    assert source.read_bytes()==before


def test_catalog_tools_adapt_to_selection_and_toggles():
    image={op.id for op in operations(['a.png'],'tools',True)}
    assert {'edit_image','annotate_image','crop_image','redact_image','frame_image'}<=image
    assert 'collage' not in image and 'collage' in {op.id for op in operations(['a.png','b.jpg'],'tools',True)}
    assert 'merge_pdf' not in {op.id for op in operations(['a.pdf'],'tools')}
    disabled={'image/crop_image','common/compress','convert/jpg'}
    for mode in ['quick','convert','tools']:
        assert not any(feature_key(['a.png'],op.id) in disabled for op in operations(['a.png'],mode,True,lambda key:key not in disabled))
    assert feature_key(['a.mp4'],'trim')!=feature_key(['a.wav'],'trim')


def test_background_image_and_collage_settings(photo,tmp_path):
    background=tmp_path/'bg.png'; Image.new('RGB',(50,50),'navy').save(background)
    engine=Engine(output_dir=tmp_path/'out')
    frame=engine.execute([photo],'frame_image',{'padding':20,'radius':10,'shadow':False,'background_image':str(background),'background_blur':4})[0]
    with Image.open(frame) as im: assert im.size==(220,160) and im.getpixel((0,0))[:3]==(0,0,128)
    collage=engine.execute([photo,background],'collage',{'columns':2,'cell':96,'gap':10,'radius':5,'background':'#112233'})[0]
    with Image.open(collage) as im: assert im.size==(222,116)


@pytest.fixture(scope='module')
def media(tmp_path_factory):
    root=tmp_path_factory.mktemp('tools-media'); engine=Engine()
    video=root/'clip.mp4'; audio=root/'audio.wav'
    engine.ff(['-f','lavfi','-i','testsrc2=s=160x120:r=10:d=2','-c:v','libx264','-pix_fmt','yuv420p',video])
    engine.ff(['-f','lavfi','-i','sine=frequency=400:duration=2','-ac','2',audio])
    return video,audio


def test_custom_video_splits_snapshots_and_blur(media,tmp_path):
    video,_=media; engine=Engine(output_dir=tmp_path/'out')
    parts=engine.execute([video],'split_media',{'points':'.5,1.2'})
    assert len(parts)==3 and sum(engine.duration(p) for p in parts)==pytest.approx(2,abs=.3)
    frames=engine.execute([video],'snapshot',{'times':'0,.8,1.5'})
    assert len(frames)==3
    for path in frames:
        with Image.open(path) as im: assert im.size==(160,120)
    blurred=engine.execute([video],'redact_video',{'x':0,'y':0,'width':80,'height':80,'start':0,'end':1,'blur':True})[0]
    assert engine.duration(blurred)==pytest.approx(2,abs=.1)
    before=tmp_path/'before.png'; after=tmp_path/'after.png'
    engine.ff(['-i',video,'-frames:v','1',before]); engine.ff(['-i',blurred,'-frames:v','1',after])
    with Image.open(before) as a,Image.open(after) as b:
        original=np.asarray(a.convert('L'),dtype=float)[:75,:75]; result=np.asarray(b.convert('L'),dtype=float)[:75,:75]
        assert np.abs(np.diff(result,axis=1)).mean()<np.abs(np.diff(original,axis=1)).mean()*.6


def test_audio_channel_gains_and_normalization_options(media,tmp_path):
    import wave
    _,audio=media; engine=Engine(output_dir=tmp_path/'out')
    stereo=engine.execute([audio],'channels',{'channels':2,'left_gain':0,'right_gain':1})[0]
    with wave.open(stereo) as wav:
        samples=np.frombuffer(wav.readframes(wav.getnframes()),dtype='<i2').reshape(-1,2)
        assert wav.getnchannels()==2 and np.max(np.abs(samples[:,0]))==0 and np.max(np.abs(samples[:,1]))>100
    normalized=engine.execute([audio],'normalize',{'loudness':-20,'true_peak':-2,'loudness_range':7})[0]
    assert engine.duration(normalized)==pytest.approx(2,abs=.1)
    with pytest.raises(ConversionError): engine.execute([audio],'normalize',{'loudness':99})
